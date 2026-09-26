"""
shared/pptx_helpers.py
Plantilla visual UCEN compartida: fondo con degradado + logo, título, bajada
(descripción de universo/población) y punteo numerado al pie.

Reemplaza el bloque de ~100 líneas que cada script de P3 copiaba y pegaba
individualmente (carga de Fondotipop.pptx, degradado, helpers de texto,
constantes de layout) — acá vive una sola vez.

CONVENCIONES (2026-08-01, pedido explícito de la contraparte) — aplican a
todos los productos:
  - Cualquier mención de un tamaño de muestra en leyendas, punteos o bajadas
    se escribe "N°=", nunca "n=" — ej. `label=f"Hombre  (N°={n_hombre})"`.
  - La bajada (`subtitulo()`) NO debe mencionar el nombre del archivo/CSV
    fuente (ej. nada de "Fuente: docentes_jornada.csv") — es un detalle interno
    de la generación, no aporta nada para la contraparte. Sí describir el
    universo/población en palabras (ej. "Universo: 624 docentes Jornada").
  - (2026-08-02) La bajada se corta después del N°/cifra principal — el detalle
    de por qué falta un dato ("...principalmente sin dotación") y la mención de
    la fuente van en `kit.notas(sl, texto)` (notas del orador), no en pantalla.
    Ej.: bajada = "...(111 sin dato)"; notas = "111 sin dato de sexo/edad,
    principalmente por falta de registro en DOTACION. Fuente: archivo.csv."

Uso típico (un script por sub-tema en products/<producto>/<carpeta>/<subtema>/):

    import sys; sys.stdout.reconfigure(encoding="utf-8")  # evita UnicodeEncodeError en consola Windows
    from pathlib import Path
    ROOT = Path(__file__).resolve().parents[4]   # ajustar según profundidad real
    sys.path.insert(0, str(ROOT))
    sys.path.insert(0, str(ROOT / "shared"))
    from config import CASCADE
    from pptx_helpers import UcenSlideKit
    from pptx import Presentation
    from pptx.util import Emu

    kit = UcenSlideKit(out_dir=Path(__file__).parent)
    kit.ensure_bg()

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(kit.SW_EMU), Emu(kit.SH_EMU)

    fig = kit.new_chart_fig()
    ax = kit.chart_axes(fig)
    # ... dibujar el gráfico en ax ...
    chart_path = kit.save_chart(fig, "mi_grafico.png")

    sl = kit.new_slide(prs)
    kit.pic(sl, prs, kit.SHARED_BG)          # fondo, a pantalla completa
    kit.pic_chart(sl, prs, chart_path)       # gráfico, en su caja real (más chico y centrado)
    kit.title(sl, "Título de la diapositiva")
    kit.subtitulo(sl, "Universo: 624 docentes Jornada  ·  513 con dato disponible")
    kit.punteo_numerado(sl, ["Primer hallazgo...", "Segundo hallazgo..."])
    prs.save(OUT_PPTX)
"""
import os
import re
import zipfile
from pathlib import Path

import numpy as np
import matplotlib.pyplot as plt
import matplotlib.text as _mtext
from PIL import Image as PILImage
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

_HERE = Path(__file__).resolve().parent
_ROOT = _HERE.parent
import sys
sys.path.insert(0, str(_ROOT))
from config import ASSETS

FONDOTIPO = os.path.join(ASSETS, "Fondotipop.pptx")


# ── Formato numérico chileno (2026-09-26, revisión de coherencia P1, obs. 15-16) ──────
# Los scripts formatean números al estilo inglés (f"{x:.1f}", f"{n:,}"). En vez de tocar
# cada f-string, todo texto que pasa por el kit (títulos, bajadas, punteos, notas, franjas)
# y todo texto de los gráficos matplotlib se convierte aquí: coma decimal, punto de miles,
# y "p=0.0000" (imposible) pasa a "p<0,0001". Regla para los scripts: escribir SIEMPRE en
# formato inglés (0.72, 134,640); nunca escribir a mano "0,72" o "134.640", porque el
# punto de miles se leería como decimal.
_RE_MILES = re.compile(r"(?<=\d),(?=\d{3}(?!\d))")
_RE_DECIMAL = re.compile(r"(?<=\d)\.(?=\d)")
_RE_P_CERO = re.compile(r"p(\s*)=(\s*)0[.,]0000\b")


def formato_cl(s):
    if not isinstance(s, str) or not any(ch.isdigit() for ch in s):
        return s
    s = _RE_P_CERO.sub(lambda m: f"p{m.group(1)}<{m.group(2)}0.0001", s)
    s = _RE_MILES.sub("\x00", s)
    s = _RE_DECIMAL.sub(",", s)
    return s.replace("\x00", ".")


def lectura_p(p, prueba="prueba t de Welch"):
    """Frase estándar para interpretar un valor p (revisión P1, obs. 5): asociación, no causa;
    'poco probable que se deba solo al azar' en vez de 'se puede descartar el azar'."""
    pv = "p<0.0001" if p < 0.0001 else f"p={p:.4f}"
    if p < 0.05:
        return (f"La {prueba} da {pv}: la diferencia es estadísticamente significativa al 5% "
                f"(es poco probable que se deba solo al azar).")
    return (f"La {prueba} da {pv}: la diferencia no es estadísticamente significativa al 5% "
            f"(podría deberse al azar).")


def prueba_dict(bloque, comparacion, resultado, n, p, prueba="t de Welch"):
    """Registro de una prueba para la tabla resumen del anexo (revisión P1, obs. 6)."""
    return dict(bloque=bloque, comparacion=comparacion, resultado=resultado, n=int(n),
                p=float(p), prueba=prueba)


if not getattr(_mtext.Text.set_text, "_formato_cl", False):
    _set_text_original = _mtext.Text.set_text

    def _set_text_cl(self, s):
        return _set_text_original(self, formato_cl(s) if isinstance(s, str) else s)

    _set_text_cl._formato_cl = True
    _mtext.Text.set_text = _set_text_cl


class UcenSlideKit:
    """Plantilla visual UCEN (fondo + título + bajada + punteo numerado).

    Una instancia por carpeta de salida (out_dir) — el fondo generado se
    cachea ahí como `_background.png` para no recalcularlo en cada corrida.
    """

    # ── Layout ──────────────────────────────────────────────────────────────
    # SW/SH en pulgadas; el resto en EMU (914400 EMU = 1 pulgada).
    # CONTENT_* = ancho de texto (título/bajada/punteo) — franja amplia, fija.
    # CHART_*   = caja del gráfico — más angosta y centrada, independiente del
    #             ancho de texto (así el gráfico se ve "más chico y centrado"
    #             sin angostar el título/bajada/punteo).
    SW, SH = 13.333, 7.5
    SW_EMU, SH_EMU = 12192000, 6858000
    _IN = 914400  # 1 pulgada en EMU

    CONTENT_L, CONTENT_W = 786581, 10599174   # franja de texto (título/bajada/punteo)
    LOGO_L, LOGO_T, LOGO_W, LOGO_H = 9813773, 656354, 1756626, 697725

    TITLE_L, TITLE_T, TITLE_W, TITLE_H = CONTENT_L, int(0.30 * _IN), CONTENT_W, int(0.50 * _IN)

    # Gap título→bajada == gap bajada→gráfico == 0.55" (misma constante _GAP_TB
    # para las dos, a pedido expreso del usuario — "mantener la distancia
    # igualitaria de título con gráfico"). Gráfico y punteo corridos más abajo
    # (2026-08-01, 2do ajuste) para usar todo el alto de la diapositiva.
    _GAP_TB = int(0.55 * _IN)
    POP_L, POP_W, POP_H = CONTENT_L, CONTENT_W, int(0.30 * _IN)
    POP_T = TITLE_T + TITLE_H + _GAP_TB

    CHART_W = int(9.5 * _IN)
    CHART_L = (SW_EMU - CHART_W) // 2                      # centrado horizontal
    CHART_T = POP_T + POP_H + _GAP_TB
    CHART_H = int(3.3 * _IN)

    BUL_L, BUL_W = CONTENT_L, CONTENT_W
    BUL_T = CHART_T + CHART_H + int(0.65 * _IN)             # gap gráfico→punteo
    BUL_H = int(1.10 * _IN)

    # Alias retrocompatibles (algunos scripts ya usan PIC_*)
    PIC_L, PIC_T, PIC_W, PIC_H = CHART_L, CHART_T, CHART_W, CHART_H

    def __init__(self, out_dir):
        self.out_dir = Path(out_dir)
        self.out_dir.mkdir(parents=True, exist_ok=True)
        self.SHARED_BG = str(self.out_dir / "_background.png")
        self._bg_arr = None
        self._logo_arr = None
        self._grad = None

    # ── Fracciones matplotlib para figuras transparentes superpuestas ─────────
    def _ex(self, e): return e / self.SW_EMU
    def _ey(self, e): return e / self.SH_EMU
    def _fig_rect(self, l, t, w, h): return (l, 1 - t - h, w, h)

    def chart_rect(self, pad=0.0):
        """Rectángulo (l, b, w, h) en fracción de figura para el área del gráfico."""
        r = self._fig_rect(self._ex(self.PIC_L), self._ey(self.PIC_T),
                            self._ex(self.PIC_W), self._ey(self.PIC_H))
        if pad:
            return (r[0] + pad, r[1] + pad, r[2] - 2 * pad, r[3] - 2 * pad)
        return r

    # ── Fondo UCEN (extrae assets de Fondotipop.pptx, cachea en out_dir) ──────
    def _load_assets(self):
        if self._bg_arr is not None:
            return
        bg_path = self.out_dir / "_fondotipo_bg.jpg"
        logo_path = self.out_dir / "_fondotipo_logo.png"
        for path, zname in [(bg_path, "ppt/media/image1.jpg"),
                             (logo_path, "ppt/media/image2.png")]:
            if not path.exists():
                with zipfile.ZipFile(FONDOTIPO) as z:
                    with open(path, "wb") as f:
                        f.write(z.read(zname))

        with PILImage.open(bg_path) as im:
            rgb = im.convert("RGB"); iw, ih = rgb.size
            nh = int(iw / (16 / 9)); y0 = min(int(ih * 0.12), ih - nh)
            self._bg_arr = np.array(rgb.crop((0, y0, iw, y0 + nh)))

        with PILImage.open(logo_path) as logo:
            self._logo_arr = np.array(logo.convert("RGBA")).astype(np.float32) / 255.0

        h_grad = 600
        grad = np.zeros((h_grad, 1, 4), dtype=np.float32)
        stops = [(0.00, (0, 33, 71)), (0.54, (0, 70, 128)), (1.00, (144, 171, 196))]
        for r in range(h_grad):
            t = r / (h_grad - 1)
            for i in range(len(stops) - 1):
                t0, c0 = stops[i]; t1, c1 = stops[i + 1]
                if t0 <= t <= t1:
                    s = (t - t0) / (t1 - t0)
                    grad[r, 0] = [(c0[0] + s * (c1[0] - c0[0])) / 255,
                                  (c0[1] + s * (c1[1] - c0[1])) / 255,
                                  (c0[2] + s * (c1[2] - c0[2])) / 255, 0.82]
                    break
        self._grad = grad

    def _bg_fig(self):
        self._load_assets()
        logo_rect = self._fig_rect(self._ex(self.LOGO_L), self._ey(self.LOGO_T),
                                    self._ex(self.LOGO_W), self._ey(self.LOGO_H))
        fig = plt.figure(figsize=(self.SW, self.SH), facecolor="#101820")
        fig.patch.set_facecolor("#101820")
        for z, arr in [(0, self._bg_arr), (1, self._grad)]:
            ax = fig.add_axes([0, 0, 1, 1], zorder=z)
            ax.imshow(arr, extent=[0, 1, 0, 1], aspect="auto", origin="upper")
            ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
        al = fig.add_axes(logo_rect, zorder=10, facecolor="none")
        al.imshow(self._logo_arr, aspect="auto"); al.axis("off"); al.patch.set_visible(False)
        return fig

    def ensure_bg(self):
        """Genera y cachea `_background.png` en out_dir si no existe todavía."""
        if not os.path.exists(self.SHARED_BG):
            fig = self._bg_fig()
            plt.savefig(self.SHARED_BG, dpi=150, facecolor=fig.get_facecolor())
            plt.close(fig)

    # ── Figuras de gráfico (transparentes, se superponen sobre el fondo) ──────
    def new_fig(self):
        """[modo legacy] Figura del tamaño completo de la diapositiva — usar
        junto con `chart_rect()` + `pic()`. Deja mucho margen transparente
        alrededor del gráfico en el PNG resultante; para un PNG ajustado al
        contenido (recomendado) usar `new_chart_fig()` + `pic_chart()`."""
        fig = plt.figure(figsize=(self.SW, self.SH), facecolor="none")
        fig.patch.set_facecolor("none")
        return fig

    def new_chart_fig(self):
        """Figura transparente del tamaño real de la caja del gráfico (CHART_W ×
        CHART_H), no de la diapositiva completa. El PNG queda ajustado al
        contenido — usar junto con `chart_axes()` + `pic_chart()`."""
        w_in = self.CHART_W / self._IN
        h_in = self.CHART_H / self._IN
        fig = plt.figure(figsize=(w_in, h_in), facecolor="none")
        fig.patch.set_facecolor("none")
        return fig

    def chart_axes(self, fig, left=0.075, right=0.02, top=0.06, bottom=0.16):
        """Ejes dentro de la figura creada con `new_chart_fig()`. Márgenes
        asimétricos por defecto — dejan espacio para xlabel/xticklabels
        (bottom) e ylabel/yticklabels (left) sin que se corten."""
        return fig.add_axes([left, bottom, 1 - left - right, 1 - top - bottom],
                             facecolor="none", zorder=5)

    def save_chart(self, fig, name):
        path = str(self.out_dir / name)
        plt.savefig(path, dpi=150, facecolor="none", transparent=True)
        plt.close(fig)
        return path

    # ── Helpers python-pptx ────────────────────────────────────────────────────
    def new_slide(self, prs):
        return prs.slides.add_slide(prs.slide_layouts[6])

    def pic(self, sl, prs, path):
        """Imagen a pantalla completa (para el fondo `SHARED_BG`)."""
        sl.shapes.add_picture(path, Emu(0), Emu(0), prs.slide_width, prs.slide_height)

    def pic_chart(self, sl, prs, chart_path):
        """Coloca el PNG del gráfico (generado con `new_chart_fig()`) en su
        posición y tamaño reales (CHART_L/T/W/H) — más chico y centrado,
        no a pantalla completa."""
        sl.shapes.add_picture(chart_path, Emu(self.CHART_L), Emu(self.CHART_T),
                               Emu(self.CHART_W), Emu(self.CHART_H))

    def txt(self, sl, text, left, top, width, height, fs=12, bold=False, italic=False,
            color="#FFFFFF", align=PP_ALIGN.LEFT, wrap=True, lspc=0, font_name="Calibri"):
        txb = sl.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
        tf = txb.text_frame; tf.word_wrap = wrap
        for i, line in enumerate(formato_cl(str(text)).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if lspc > 0 and i > 0:
                p.space_before = Pt(lspc)
            run = p.add_run(); run.text = line
            run.font.size = Pt(fs); run.font.bold = bold; run.font.italic = italic
            if font_name:
                run.font.name = font_name
            r, g, b = int(color[1:3], 16), int(color[3:5], 16), int(color[5:7], 16)
            run.font.color.rgb = RGBColor(r, g, b)

    def title(self, sl, text, fs=20):
        """Título de la diapositiva."""
        self.txt(sl, text, self.TITLE_L, self.TITLE_T, self.TITLE_W, self.TITLE_H,
                  fs=fs, bold=True, color="#FFFFFF", align=PP_ALIGN.CENTER)

    def subtitulo(self, sl, text):
        """Bajada explicativa de universo/población, bajo el título.
        No mencionar el nombre del archivo/CSV fuente (ej. "Fuente: x.csv") —
        es un detalle interno, no aporta nada para la contraparte."""
        self.txt(sl, text, self.POP_L, self.POP_T, self.POP_W, self.POP_H,
                  fs=10, italic=True, color="#C8DCF0", font_name="Calibri")

    def punteo_numerado(self, sl, items, fs=13):
        """Lista numerada (1. 2. ...) al pie de la diapositiva."""
        txb = sl.shapes.add_textbox(Emu(self.BUL_L), Emu(self.BUL_T),
                                     Emu(self.BUL_W), Emu(self.BUL_H))
        tf = txb.text_frame; tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(6); p.alignment = PP_ALIGN.LEFT
            run = p.add_run(); run.text = f"{i + 1}.  {formato_cl(item)}"
            run.font.size = Pt(fs)
            run.font.name = "Calibri"
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    def notas(self, sl, text):
        """Notas del orador (no se ven en la diapositiva, solo en modo presentador/al exportar).
        Uso: mover ahí detalle secundario que sobra en la bajada visible (ej. la explicación
        larga de por qué falta un dato, dejando en pantalla solo el N°)."""
        sl.notes_slide.notes_text_frame.text = formato_cl(text)

    # ── Layouts estructurales (portada, grillas de bloque, universo/índice/hallazgos) ──
    def portada(self, sl, titulo, subtitulo1, footer_lines, subtitulo2=None):
        """Diapositiva de portada: título + bajada(s) + pie de 1 o más líneas, centrados.
        El caller ya debe haber puesto el fondo (`kit.pic(sl, prs, kit.SHARED_BG)`)."""
        self.txt(sl, titulo, self.CONTENT_L, int(2.35 * self._IN), self.CONTENT_W, int(0.65 * self._IN),
                  fs=30, bold=True, color="#FFFFFF", align=PP_ALIGN.CENTER)
        self.txt(sl, subtitulo1, self.CONTENT_L, int(3.20 * self._IN), self.CONTENT_W, int(0.45 * self._IN),
                  fs=17, color="#C8DCF0", align=PP_ALIGN.CENTER)
        if subtitulo2:
            self.txt(sl, subtitulo2, self.CONTENT_L, int(3.72 * self._IN), self.CONTENT_W, int(0.40 * self._IN),
                      fs=12, italic=True, color="#C8DCF0", align=PP_ALIGN.CENTER)
        self.txt(sl, "\n".join(footer_lines), self.CONTENT_L, int(5.55 * self._IN),
                  self.CONTENT_W, int(1.1 * self._IN), fs=13, bold=True, color="#FFFFFF",
                  align=PP_ALIGN.CENTER, lspc=4)

    def _caja_navy(self, sl, l, t, w, h, header, body, fs_header=11, fs_body=9):
        sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(l), Emu(t), Emu(w), Emu(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0, 33, 71)
        sh.line.fill.background()
        self.txt(sl, header, l + 40000, t + 38000, w - 80000, 300000,
                  fs=fs_header, bold=True, color="#90ABC4")
        self.txt(sl, body, l + 40000, t + 360000, w - 80000, h - 420000,
                  fs=fs_body, color="#FFFFFF", lspc=3)

    def caja_grid_2x2(self, sl, cajas):
        """4 cajas navy en grilla 2×2, ocupando toda la franja de contenido (bajo el título/bajada).
        `cajas` = lista de hasta 4 tuplas (header, body)."""
        box_t = self.CHART_T
        box_b = self.BUL_T + self.BUL_H
        pad_x, gap_x, gap_y, pad_t, pad_b = 100000, 80000, 50000, 20000, 20000
        bw = (self.CONTENT_W - 2 * pad_x - gap_x) // 2
        bh = (box_b - box_t - pad_t - gap_y - pad_b) // 2
        pos = [(self.CONTENT_L + pad_x, box_t + pad_t),
               (self.CONTENT_L + pad_x + bw + gap_x, box_t + pad_t),
               (self.CONTENT_L + pad_x, box_t + pad_t + bh + gap_y),
               (self.CONTENT_L + pad_x + bw + gap_x, box_t + pad_t + bh + gap_y)]
        for (bx, by), (hdr, body) in zip(pos, cajas):
            self._caja_navy(sl, bx, by, bw, bh, hdr, body)

    def caja_grid_1x3(self, sl, cajas):
        """3 cajas navy en una sola fila, ocupando toda la franja de contenido — para
        comparar 3 grupos/categorías lado a lado (ej. 3 tercios de dificultad).
        `cajas` = lista de hasta 3 tuplas (header, body)."""
        box_t = self.CHART_T
        box_b = self.BUL_T + self.BUL_H
        pad_x, gap_x, pad_t, pad_b = 100000, 60000, 20000, 20000
        bw = (self.CONTENT_W - 2 * pad_x - 2 * gap_x) // 3
        bh = box_b - box_t - pad_t - pad_b
        pos = [(self.CONTENT_L + pad_x + i * (bw + gap_x), box_t + pad_t) for i in range(3)]
        for (bx, by), (hdr, body) in zip(pos, cajas):
            self._caja_navy(sl, bx, by, bw, bh, hdr, body, fs_header=12, fs_body=10)

    def franjas(self, sl, filas, top=None, fs=9, col_split=0.41):
        """Formato "franjas" (pedido de la contraparte 2026-09-25): filas horizontales de
        ancho completo, sin cajas de fondo, letras blancas, separadas por líneas finas.
        Cada fila = (header, columnas): header a la izquierda (22% del ancho) y 1 o 2
        columnas de texto a la derecha (`columnas` = lista de 1 o 2 strings; con 2, la
        primera ocupa `col_split` del ancho total). El alto de cada fila es proporcional
        a sus líneas estimadas, para que la más cargada no se desborde."""
        top = top if top is not None else self.POP_T + self.POP_H + int(0.12 * self._IN)
        bottom = self.BUL_T + self.BUL_H
        l, w = self.CONTENT_L, self.CONTENT_W
        col_hdr = int(w * 0.22)
        pad = 60000

        def _anchos(cols):
            if len(cols) == 1:
                return [w - col_hdr]
            a = int(w * col_split)
            return [a, w - col_hdr - a]

        def _lineas(text, ancho_emu):
            chars = max(20, int(ancho_emu / self._IN * 17.5 * 9 / fs))
            return sum(max(1, -(-len(p) // chars)) for p in str(text).split("\n"))

        pesos = []
        for hdr, cols in filas:
            n = max([_lineas(c, a) for c, a in zip(cols, _anchos(cols))] + [_lineas(hdr, col_hdr) * 1.4])
            pesos.append(n + 1)
        alturas = [(bottom - top) * p // sum(pesos) for p in pesos]

        def _linea(y):
            ln = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(l), Emu(y), Emu(w), Emu(9000))
            ln.fill.solid(); ln.fill.fore_color.rgb = RGBColor(90, 140, 190)
            ln.line.fill.background()

        _linea(top)
        y = top
        for (hdr, cols), fh in zip(filas, alturas):
            self.txt(sl, hdr, l, y + pad, col_hdr - pad, fh - 2 * pad,
                      fs=12, bold=True, color="#FFFFFF")
            x = l + col_hdr
            for c, a in zip(cols, _anchos(cols)):
                if c:
                    self.txt(sl, c, x, y + pad, a - pad, fh - 2 * pad, fs=fs, color="#FFFFFF", lspc=2)
                x += a
            y += fh
            _linea(y)

    def franjas_bloques(self, sl, franjas, top=None):
        """Grilla de apertura en franjas: una fila por bloque, con
        nombre del bloque | descripción (viñetas) | "Qué vas a ver" (lista numerada).
        `franjas` = lista de tuplas (header, descripcion_items, que_vas_a_ver_items)."""
        filas = []
        for hdr, desc, ver in franjas:
            c1 = "\n".join(f"•  {d}" for d in desc)
            c2 = ("Qué vas a ver\n" + "\n".join(f"{k + 1}. {v}" for k, v in enumerate(ver))) if ver else ""
            filas.append((hdr, [c1, c2]))
        self.franjas(sl, filas, top=top)

    def franjas_universo_indice_hallazgos(self, sl, universo_txt, indice_items, hallazgos_items):
        """Universo / Índice / Hallazgos en formato franjas (reemplaza a la caja navy de
        3 columnas de `caja_universo_indice_hallazgos`). Un índice largo (>5 ítems) se
        reparte en 2 columnas para no alargar la fila."""
        numerados = [f"{i + 1}. {it}" for i, it in enumerate(indice_items)]
        if len(numerados) > 5:
            mitad = -(-len(numerados) // 2)
            indice_cols = ["\n".join(numerados[:mitad]), "\n".join(numerados[mitad:])]
        else:
            indice_cols = ["\n".join(numerados)]
        self.franjas(sl, [
            ("Universo", [universo_txt]),
            ("Índice", indice_cols),
            ("Hallazgos", ["\n".join(f"{i + 1}. {it}" for i, it in enumerate(hallazgos_items))]),
        ], fs=10, col_split=0.39)

    def tabla(self, sl, encabezados, filas, anchos, top=None, fs=9, resaltar=None):
        """Tabla simple en formato franjas (sin cajas de color): encabezado en negrita,
        filas en blanco, línea fina entre filas. `anchos` = fracciones del ancho de contenido.
        `resaltar(i)` -> True pinta la fila i en dorado (ej. pruebas significativas)."""
        top = top if top is not None else self.POP_T + self.POP_H + int(0.12 * self._IN)
        bottom = self.BUL_T + self.BUL_H
        l, w = self.CONTENT_L, self.CONTENT_W
        n = len(filas) + 1
        fh = (bottom - top) // n
        xs = [l + int(w * sum(anchos[:j])) for j in range(len(anchos))]
        ws = [int(w * a) for a in anchos]

        def _linea(y, alpha_rgb=(90, 140, 190)):
            ln = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(l), Emu(y), Emu(w), Emu(6000))
            ln.fill.solid(); ln.fill.fore_color.rgb = RGBColor(*alpha_rgb)
            ln.line.fill.background()

        for x, cw, h in zip(xs, ws, encabezados):
            self.txt(sl, h, x, top, cw - 30000, fh, fs=fs, bold=True, color="#FFFFFF")
        _linea(top + fh)
        for i, fila in enumerate(filas):
            y = top + (i + 1) * fh
            color = "#F2D675" if (resaltar and resaltar(i)) else "#FFFFFF"
            for x, cw, v in zip(xs, ws, fila):
                self.txt(sl, str(v), x, y + 15000, cw - 30000, fh, fs=fs, color=color)
            _linea(y + fh, (60, 90, 125))

    def caja_universo_indice_hallazgos(self, sl, universo_txt, indice_items, hallazgos_items):
        """1 caja navy de ancho completo, 3 columnas: Universo (párrafo) | Índice (lista) | Hallazgos (lista)."""
        box_l, box_t = self.CONTENT_L, self.CHART_T
        box_w = self.CONTENT_W
        box_h = (self.BUL_T + self.BUL_H) - box_t
        sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(box_l), Emu(box_t), Emu(box_w), Emu(box_h))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0, 33, 71)
        sh.line.fill.background()

        col_w = box_w // 3
        for i, h in enumerate(["Universo", "Índice", "Hallazgos"]):
            cx = box_l + i * col_w
            self.txt(sl, h, cx + 50000, box_t + 45000, col_w - 100000, 260000,
                      fs=13, bold=True, color="#FFFFFF", align=PP_ALIGN.CENTER)
            div = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(cx + 50000), Emu(box_t + 330000),
                                       Emu(col_w - 100000), Emu(9000))
            div.fill.solid(); div.fill.fore_color.rgb = RGBColor(90, 140, 190)
            div.line.fill.background()

        self.txt(sl, universo_txt, box_l + 50000, box_t + 410000, col_w - 100000, box_h - 470000,
                  fs=10, color="#FFFFFF", lspc=4)
        idx_txt = "\n".join(f"{i + 1}. {it}" for i, it in enumerate(indice_items))
        self.txt(sl, idx_txt, box_l + col_w + 50000, box_t + 410000, col_w - 100000, box_h - 470000,
                  fs=9.5, color="#FFFFFF", lspc=5)
        hal_txt = "\n".join(f"{i + 1}. {it}" for i, it in enumerate(hallazgos_items))
        self.txt(sl, hal_txt, box_l + 2 * col_w + 50000, box_t + 410000, col_w - 100000, box_h - 470000,
                  fs=9.5, color="#FFFFFF", lspc=5)
