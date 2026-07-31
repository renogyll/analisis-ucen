"""
Standalone: DIAPO_metodologia.pptx  — 1 slide
Anexo metodológico con 2 columnas y 2 títulos de sección.
 - Izq:  Universo de Análisis (embudo + tipos, desde datos reales)
 - Der:  Marco Metodológico   (z-score + criterios de exclusión + métricas)
Sin logo doble: el fondo ya lo incluye.
"""
import sys; sys.stdout.reconfigure(encoding="utf-8")
import os, zipfile, pathlib
import pandas as pd
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Rutas ─────────────────────────────────────────────────────────────────────
BASE    = os.path.dirname(os.path.abspath(__file__))
REPO    = str(pathlib.Path(BASE).parents[1])
CASCADE = os.path.join(REPO, "data", "cascade")
OUT_DIR = os.path.join(BASE, "dark_slides_v3")
SHARED_BG = os.path.join(OUT_DIR, "_background.png")  # ya contiene el logo
OUT_PPTX  = r"c:\Users\r.gonzalez_fluxsolar.LAPTOP-FLUX-ECO\Downloads\DIAPO_metodologia_v3.pptx"

# ── Layout (igual que generar_presentacion.py) ────────────────────────────────
SW_EMU = 12192000
SH_EMU = 6858000

# ── Datos reales ──────────────────────────────────────────────────────────────
sat = pd.read_csv(os.path.join(CASCADE, "05_aptos_p3", "p3_sat_zscore.csv"),
                  encoding="utf-8-sig")
vc  = sat["tipos_formacion"].value_counts()
n_t = len(sat)  # 316

lbl = [str(t).title().replace("Taller", "Oferta formativa") for t in vc.index]
val = vc.values.tolist()

n_pura_of = sum(v for l, v in zip(lbl, val) if l == "Oferta formativa")
n_pura_di = sum(v for l, v in zip(lbl, val) if l == "Diplomado")
n_pura_pr = sum(v for l, v in zip(lbl, val) if l == "Proyecto")
n_mixta   = sum(v for l, v in zip(lbl, val) if "|" in l)

print(f"Universo: {n_t}  |  OF pura={n_pura_of}  Di pura={n_pura_di}  Pr pura={n_pura_pr}  Mixta={n_mixta}")

# ── Colores ───────────────────────────────────────────────────────────────────
W   = RGBColor(0xFF, 0xFF, 0xFF)   # blanco
CEL = RGBColor(0x7E, 0xBF, 0xF0)   # celeste encabezado
AMA = RGBColor(0xFF, 0xC8, 0x00)   # amarillo fórmula
GRI = RGBColor(0xCC, 0xDD, 0xEE)   # gris claro POP

# ── PPTX ─────────────────────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width  = Emu(SW_EMU)
prs.slide_height = Emu(SH_EMU)
sl = prs.slides.add_slide(prs.slide_layouts[6])

# Fondo (sin logo adicional — ya está en el PNG)
sl.shapes.add_picture(SHARED_BG, Emu(0), Emu(0), Emu(SW_EMU), Emu(SH_EMU))

# ── Helper de texto ───────────────────────────────────────────────────────────
def box(text, lx, ty, w, h,
        sz=10.5, bold=False, italic=False,
        color=W, align=PP_ALIGN.LEFT):
    txb = sl.shapes.add_textbox(Emu(lx), Emu(ty), Emu(w), Emu(h))
    tf  = txb.text_frame; tf.word_wrap = True
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run(); run.text = line
        run.font.size    = Pt(sz)
        run.font.bold    = bold
        run.font.italic  = italic
        run.font.color.rgb = color
        run.font.name    = "Calibri"

def mbox(lines, lx, ty, w, h, default_sz=10, default_color=W, gap=40000):
    """Multi-line box: cada elemento es (texto, sz, bold, color) o str."""
    txb = sl.shapes.add_textbox(Emu(lx), Emu(ty), Emu(w), Emu(h))
    tf  = txb.text_frame; tf.word_wrap = True
    first = True
    for item in lines:
        if isinstance(item, str):
            txt, sz, bld, col = item, default_sz, False, default_color
        else:
            txt = item[0]
            sz  = item[1] if len(item) > 1 else default_sz
            bld = item[2] if len(item) > 2 else False
            col = item[3] if len(item) > 3 else default_color
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.space_after = Emu(gap)
        run = p.add_run(); run.text = txt
        run.font.size  = Pt(sz)
        run.font.bold  = bld
        run.font.color.rgb = col
        run.font.name  = "Calibri"
        first = False

# ═══════════════════════════════════════════════════════════════════════════════
# TÍTULO PRINCIPAL
# ═══════════════════════════════════════════════════════════════════════════════
box("Anexo Metodológico — Universo y Marco de Análisis P3",
    lx=786581, ty=170000, w=10599174, h=650000,
    sz=19, bold=True, align=PP_ALIGN.CENTER)

box("Programa P3  ·  1.144 docentes UCEN  ·  Períodos 2022–2025  ·  Análisis antes–durante–después",
    lx=786581, ty=820000, w=10200000, h=260000,
    sz=8.5, italic=True, color=GRI, align=PP_ALIGN.CENTER)

# ── Línea divisoria vertical ─────────────────────────────────────────────────
div_x = 5350000
from pptx.util import Inches
from lxml import etree
# Dibujamos un rectángulo delgado como divisor
div_shape = sl.shapes.add_shape(1,
    Emu(div_x), Emu(1090000), Emu(18000), Emu(5550000))
div_shape.fill.solid()
div_shape.fill.fore_color.rgb = RGBColor(0x3A, 0x6E, 0xA8)
div_shape.line.fill.background()

# ═══════════════════════════════════════════════════════════════════════════════
# COLUMNA IZQUIERDA — Universo de Análisis
# ═══════════════════════════════════════════════════════════════════════════════
LX = 786581
LW = 4380000
COL_HDR_H = 320000

# Encabezado izquierdo
box("Universo de Análisis",
    lx=LX, ty=1100000, w=LW, h=COL_HDR_H,
    sz=12, bold=True, color=CEL)

# Embudo
mbox([
    ("1.144  docentes  (universo base)", 13, True, W),
    ("Jornada + Honorario — todos los períodos 2022–2025", 8.5, False, GRI),
    ("", 4, False, W),
    ("▼   726  con ≥1 instancia de formación P3", 11, True, W),
    ("Participaron en Oferta formativa, Diplomado o Proyecto", 8.5, False, GRI),
    ("", 4, False, W),
    ("▼   316  Aptos P3  —  universo de análisis central", 12, True, W),
    ("SAT válido ≥3 períodos consecutivos (pre · durante · post)", 8.5, False, GRI),
    ("", 4, False, W),
    ("▼   206  Aptos P3 con EDD  —  Bloque IV", 10.5, True, W),
    ("110 de los 316 no poseen registro EDD", 8.5, False, GRI),
],
    lx=LX, ty=1450000, w=LW, h=2500000, gap=20000)

# Separador visual
sep1 = sl.shapes.add_shape(1, Emu(LX), Emu(3960000), Emu(LW), Emu(12000))
sep1.fill.solid(); sep1.fill.fore_color.rgb = RGBColor(0x3A, 0x6E, 0xA8)
sep1.line.fill.background()

# Tipos de oferta formativa (sin título propio, introducidos con texto)
mbox([
    ("Tipos de Oferta Formativa  (de los 316 Aptos P3)", 10.5, True, CEL),
    ("", 3, False, W),
    (f"Oferta formativa  (actividades, 1 semestre)  —  N°={n_pura_of} docentes puros", 9.5, False, W),
    (f"Diplomado  (larga duración, ≥2 semestres)     —  N°={n_pura_di} docentes puros", 9.5, False, W),
    (f"Proyecto de Innovación Docente                 —  N°={n_pura_pr} docentes puros", 9.5, False, W),
    (f"Participación Mixta  (combinación de tipos)   —  N°={n_mixta} docentes", 9.5, False, W),
],
    lx=LX, ty=4000000, w=LW, h=1600000, gap=22000)

# ═══════════════════════════════════════════════════════════════════════════════
# COLUMNA DERECHA — Marco Metodológico
# ═══════════════════════════════════════════════════════════════════════════════
RX = 5500000
RW = 5800000

# Encabezado derecho
box("Marco Metodológico",
    lx=RX, ty=1100000, w=RW, h=COL_HDR_H,
    sz=12, bold=True, color=CEL)

# Fórmula z-score
box("z  =  (SAT docente  −  media facultad · período)  /  DE facultad · período",
    lx=RX, ty=1440000, w=RW, h=340000,
    sz=11, bold=True, color=AMA)

# Interpretación
mbox([
    ("z = 0  →  promedio exacto de la facultad ese semestre", 9.5, False, W),
    ("z > 0  →  sobre el promedio de la facultad", 9.5, False, W),
    ("z < 0  →  bajo el promedio de la facultad", 9.5, False, W),
    ("Permite comparar docentes entre facultades con distintos niveles de exigencia.", 8.5, False, GRI),
],
    lx=RX, ty=1790000, w=RW, h=780000, gap=16000)

# Separador
sep3 = sl.shapes.add_shape(1, Emu(RX), Emu(2580000), Emu(RW), Emu(12000))
sep3.fill.solid(); sep3.fill.fore_color.rgb = RGBColor(0x3A, 0x6E, 0xA8)
sep3.line.fill.background()

# Criterios de exclusión
box("Criterios de exclusión  (aplican a todos los bloques)",
    lx=RX, ty=2620000, w=RW, h=290000,
    sz=10.5, bold=True, color=W)

mbox([
    ("•  Cobertura ≥40%: sección excluida si < 40% de alumnos respondió la encuesta.", 9.5, False, W),
    ("•  Sección mínima: se excluyen secciones con < 7 alumnos matriculados.", 9.5, False, W),
    ("•  Asignaturas excluidas: práctica profesional, proyecto de título, seminario,", 9.5, False, W),
    ("    ciclo formativo e integración profesional (escala de logro no comparable).", 9.5, False, W),
    ("•  Período de referencia pre: promedio de todos los semestres anteriores a la", 9.5, False, W),
    ("    primera instancia formativa del docente.", 9.5, False, W),
    ("•  Evolución temporal: docentes con datos en < 2 períodos se excluyen.", 9.5, False, W),
],
    lx=RX, ty=2920000, w=RW, h=1700000, gap=18000)

# Separador
sep4 = sl.shapes.add_shape(1, Emu(RX), Emu(4630000), Emu(RW), Emu(12000))
sep4.fill.solid(); sep4.fill.fore_color.rgb = RGBColor(0x3A, 0x6E, 0xA8)
sep4.line.fill.background()

# Métricas de evaluación (movidas aquí desde columna izquierda)
box("Métricas de evaluación principal",
    lx=RX, ty=4670000, w=RW, h=270000,
    sz=10.5, bold=True, color=W)

mbox([
    ("•  SAT Nota (1–7) estandarizada como z-score  ·  SAT % Recomendación (dicotómica Sí/No)", 9.5, False, W),
    ("•  EDD: Evaluación de Desempeño Docente — escala 0–1, expresada en %", 9.5, False, W),
    ("•  Rendimiento Académico: nota promedio alumnos (1–7)  ·  % Aprobación ≥4.0", 9.5, False, W),
],
    lx=RX, ty=4950000, w=RW, h=650000, gap=18000)

# Separador
sep5 = sl.shapes.add_shape(1, Emu(RX), Emu(5610000), Emu(RW), Emu(12000))
sep5.fill.solid(); sep5.fill.fore_color.rgb = RGBColor(0x3A, 0x6E, 0xA8)
sep5.line.fill.background()

# Grupos por bloque (compacto al pie)
mbox([
    ("Grupos de análisis  ·  "
     f"B-I {n_t} Aptos P3  ·  "
     f"B-II {n_t} trat / 300 ctrl  ·  "
     "B-III 316 (perfil) / 619 (notas) / 300 ctrl  ·  "
     "B-IV 206 trat / 130 ctrl", 8.5, False, GRI),
],
    lx=RX, ty=5640000, w=RW, h=600000, gap=0)

prs.save(OUT_PPTX)
print(f"✓ Guardado: {OUT_PPTX}")
