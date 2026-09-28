"""
Standalone: genera DIAPO_combinaciones.pptx (1 slide)
Gráfico de barras horizontales — Combinaciones de Modalidad de Formación.
Replica exactamente el estilo y layout de generar_presentacion.py.
"""
import sys; sys.stdout.reconfigure(encoding="utf-8")
import os, zipfile, pathlib
import numpy as np
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patheffects as pe
import pandas as pd
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Rutas ─────────────────────────────────────────────────────────────────────
BASE    = os.path.dirname(os.path.abspath(__file__))
REPO    = str(pathlib.Path(BASE).parents[3])
CASCADE = os.path.join(REPO, "data", "cascade")
SCRATCH = os.path.join(REPO, "outputs", "scratch")
OUT_DIR = os.path.join(BASE, "dark_slides_v3")
FONDOTIPO = os.path.join(REPO, "assets", "Fondotipop.pptx")

BG_PATH   = os.path.join(SCRATCH, "fondotipo_image1.jpg")
LOGO_PATH = os.path.join(SCRATCH, "fondotipo_image2.png")
SHARED_BG = os.path.join(OUT_DIR, "_background.png")
OUT_PPTX  = os.path.join(REPO, "outputs", "pptx", "DIAPO_combinaciones.pptx")

# Extraer assets si no existen
for path, zname in [(BG_PATH,"ppt/media/image1.jpg"),(LOGO_PATH,"ppt/media/image2.png")]:
    if not os.path.exists(path):
        with zipfile.ZipFile(FONDOTIPO) as z:
            with open(path,"wb") as f: f.write(z.read(zname))

# ── Constantes de layout (idénticas a generar_presentacion.py) ───────────────
SW, SH       = 13.333, 7.5
SW_EMU       = 12192000
SH_EMU       = 6858000

PIC_L, PIC_T, PIC_W, PIC_H = 786581, 1125000, 10599174, 3720000
BUL_L, BUL_T, BUL_W, BUL_H = 786581, 4870000, 10599174, 1870000
LOGO_L, LOGO_T, LOGO_W, LOGO_H = 9813773, 656354, 1756626, 697725
TITLE_L, TITLE_T, TITLE_W, TITLE_H = PIC_L, 185000, PIC_W, 710000
POP_L,   POP_T,   POP_W,   POP_H   = PIC_L, 845000, 9000000, 255000

def _ex(e): return e / SW_EMU
def _ey(e): return e / SH_EMU
def _fig_rect(l, t, w, h): return (l, 1 - t - h, w, h)

PIC_RECT = _fig_rect(_ex(PIC_L), _ey(PIC_T), _ex(PIC_W), _ey(PIC_H))
CHART_X  = PIC_RECT[0] + 0.13
CHART_Y  = PIC_RECT[1] + 0.04
CHART_W  = PIC_RECT[2] - 0.19
CHART_H  = PIC_RECT[3] - 0.09

CTITLE_L = int(CHART_X * SW_EMU)
CTITLE_T = PIC_T + 38000
CTITLE_W = int((PIC_RECT[0] + PIC_RECT[2] - CHART_X + 0.02) * SW_EMU)
CTITLE_H = 295000

PAL = ["#5C9BD6","#64B5F6","#80DEEA","#A5D6A7","#FFB74D","#CE93D8","#90A4AE","#F48FB1"]

# ── Fondo compartido ──────────────────────────────────────────────────────────
if not os.path.exists(SHARED_BG):
    print("Generando fondo...")
    with PILImage.open(BG_PATH) as im:
        rgb = im.convert("RGB"); iw, ih = rgb.size
        nh  = int(iw / (16/9)); y0 = min(int(ih*0.12), ih - nh)
        bg_arr = np.array(rgb.crop((0, y0, iw, y0 + nh)))
    with PILImage.open(LOGO_PATH) as lg:
        logo_arr = np.array(lg.convert("RGBA")).astype(np.float32) / 255.0
    H_GRAD = 600; grad = np.zeros((H_GRAD, 1, 4), dtype=np.float32)
    for r in range(H_GRAD):
        t = r / (H_GRAD - 1)
        stops = [(0.00,(0,33,71)),(0.54,(0,70,128)),(1.00,(144,171,196))]
        for i in range(len(stops)-1):
            t0,c0 = stops[i]; t1,c1 = stops[i+1]
            if t0 <= t <= t1:
                s = (t-t0)/(t1-t0)
                grad[r,0] = [(c0[j]+s*(c1[j]-c0[j]))/255 for j in range(3)] + [0.82]; break
    fig_bg = plt.figure(figsize=(SW, SH)); ax_bg = fig_bg.add_axes([0,0,1,1])
    ax_bg.imshow(bg_arr, aspect="auto", extent=[0,1,0,1])
    ax_bg.imshow(np.tile(grad, (1, 1000, 1)), aspect="auto", extent=[0,1,0,1])
    lh, lw = logo_arr.shape[:2]
    lx = _ex(LOGO_L); ly = 1 - _ey(LOGO_T) - _ey(LOGO_H)
    ax_bg.imshow(logo_arr, aspect="auto",
                 extent=[lx, lx+_ex(LOGO_W), ly, ly+_ey(LOGO_H)])
    ax_bg.axis("off")
    plt.savefig(SHARED_BG, dpi=150, facecolor=fig_bg.get_facecolor())
    plt.close()
    print(f"  Fondo guardado: {SHARED_BG}")

# ── Cargar datos ──────────────────────────────────────────────────────────────
sat = pd.read_csv(os.path.join(CASCADE, "05_aptos_p3", "p3_sat_zscore.csv"),
                  encoding="utf-8-sig")
vc  = sat["tipos_formacion"].value_counts()
n_t = len(sat)

lbl = [str(t).title().replace("Taller", "Oferta formativa") for t in vc.index.tolist()]
val = vc.values.tolist()
pct = [100 * v / n_t for v in val]

n_solo  = sum(v for l, v in zip(lbl, val) if "|" not in l)
n_multi = n_t - n_solo

print("Labels:")
for l, v, p in zip(lbl, val, pct):
    print(f"  {l!r:45s}  {v}  ({p:.1f}%)")

# ── Gráfico (replica exacta de _hbar) ────────────────────────────────────────
fig = plt.figure(figsize=(SW, SH), facecolor="none")
fig.patch.set_facecolor("none")

n  = len(lbl); yp = np.arange(n)
ax = fig.add_axes([CHART_X, CHART_Y, CHART_W, CHART_H], facecolor="none", zorder=5)

ax.barh(yp[::-1], val, color=PAL[:n], height=0.58, edgecolor="none", alpha=0.90)

mv = max(val)
for i, (v, p) in enumerate(zip(val[::-1], pct[::-1])):
    ax.text(v + mv * 0.020, i, f"{v}  ({p:.1f}%)",
            va="center", ha="left", fontsize=10.5, fontweight="bold", color="white",
            path_effects=[pe.withStroke(linewidth=2.5, foreground="#0A0F18")])

ax.set_yticks(yp)
ax.set_yticklabels(lbl[::-1], fontsize=10.5, fontweight="bold", color="white")
ax.tick_params(axis="y", length=0, pad=8)
ax.tick_params(axis="x", colors="#AAAAAA", labelsize=9)
for sp in ax.spines.values():
    sp.set_edgecolor("white"); sp.set_alpha(0.35); sp.set_linewidth(0.9)
ax.set_xlim(0, mv * 1.55)
ax.set_ylim(-0.5, n - 0.5)
ax.xaxis.grid(True, color="white", alpha=0.07, linewidth=0.5)
ax.set_axisbelow(True)

chart_path = os.path.join(OUT_DIR, "combinaciones_chart.png")
fig.savefig(chart_path, dpi=150, facecolor="none", transparent=True)
plt.close()
print(f"\nChart guardado: {chart_path}")

# ── Armar slide ───────────────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width  = Emu(SW_EMU)
prs.slide_height = Emu(SH_EMU)
sl = prs.slides.add_slide(prs.slide_layouts[6])

# Fondo + chart
sl.shapes.add_picture(SHARED_BG, Emu(0), Emu(0), Emu(SW_EMU), Emu(SH_EMU))
sl.shapes.add_picture(chart_path, Emu(0), Emu(0), Emu(SW_EMU), Emu(SH_EMU))

def _txt(text, left, top, width, height,
         fs=12, bold=False, italic=False, color="#FFFFFF",
         align=PP_ALIGN.LEFT, wrap=True):
    txb = sl.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf  = txb.text_frame; tf.word_wrap = wrap
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run(); run.text = line
        run.font.size = Pt(fs); run.font.bold = bold; run.font.italic = italic
        r,g,b = int(color[1:3],16), int(color[3:5],16), int(color[5:7],16)
        run.font.color.rgb = RGBColor(r,g,b)

# Título
_txt("Combinaciones de Modalidad de Formación",
     TITLE_L, TITLE_T, TITLE_W, TITLE_H,
     fs=20, bold=True, color="#FFFFFF", align=PP_ALIGN.CENTER)

# POP subtitle
_txt(f"Universo: {n_t} Aptos P3  ·  todos formados  ·  Periodos 2022–2025",
     POP_L, POP_T, POP_W, POP_H,
     fs=7.5, italic=True, color="#C8DCF0")

# Chart title
_txt(f"Combinaciones de tipos de formación — Formados Aptos P3  (nº {n_t})",
     CTITLE_L, CTITLE_T, CTITLE_W, CTITLE_H,
     fs=9, color="#FFFFFF")

# Bullets
bul_items = [
    f"{n_solo} de los {n_t} formados ({100*n_solo/n_t:.0f}%) participaron en una sola modalidad; "
    f"{n_multi} tienen Participación Mixta (combinaron dos o más tipos de formación).",
    "Las combinaciones con Participación Mixta (Diplomado | Oferta formativa, Proyecto | Oferta formativa) "
    "representan el subconjunto con mayor exposición acumulada a formación.",
]
txb = sl.shapes.add_textbox(Emu(BUL_L), Emu(BUL_T), Emu(BUL_W), Emu(BUL_H))
tf  = txb.text_frame; tf.word_wrap = True
for i, item in enumerate(bul_items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(5); p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = f"•  {item}"
    run.font.size = Pt(11.5)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

prs.save(OUT_PPTX)
print(f"\n✓ Guardado: {OUT_PPTX}")
