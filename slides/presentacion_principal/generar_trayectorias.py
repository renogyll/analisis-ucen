"""
generar_trayectorias.py
3 slides de trayectorias individuales a través del tiempo:
  Slide 1 — B2: SAT individual de los 210 Aptos P3 por período
  Slide 2 — B3: Nota promedio alumnos (formados vs no formados) por período
  Slide 3 — B4: EDD jefaturas (formados vs no formados) por año
Ejecutar desde analisis-ucen/ o directamente.
"""
import sys; sys.stdout.reconfigure(encoding="utf-8")
import os, zipfile
from pathlib import Path
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patheffects as pe
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from config import CASCADE

# ── Paths ──────────────────────────────────────────────────────────────────────
REPO     = Path(__file__).resolve().parents[2]
SCRATCH  = os.path.join(REPO, "outputs", "scratch")
OUT_DIR  = os.path.join(REPO, "outputs", "scratch", "tray_slides")
OUT_PPTX = os.path.join(REPO, "outputs", "pptx", "TRAYECTORIAS_v1.pptx")
FONDOTIPO = os.path.join(REPO, "PRESENTACION_800.pptx")
os.makedirs(OUT_DIR, exist_ok=True)
os.makedirs(os.path.join(REPO, "outputs", "pptx"), exist_ok=True)

COMP = os.path.join(CASCADE, "complementarios")

# ── Assets: fondo + logo ───────────────────────────────────────────────────────
BG_PATH   = os.path.join(SCRATCH, "fondotipo_image1.jpg")
LOGO_PATH = os.path.join(SCRATCH, "fondotipo_image2.png")
for path, zname in [(BG_PATH, "ppt/media/image1.jpg"), (LOGO_PATH, "ppt/media/image2.png")]:
    if not os.path.exists(path):
        with zipfile.ZipFile(FONDOTIPO) as z:
            with open(path, "wb") as f:
                f.write(z.read(zname))

with PILImage.open(BG_PATH) as _im:
    _rgb = _im.convert("RGB"); _iw, _ih = _rgb.size
    _nh  = int(_iw / (16/9)); _y0 = min(int(_ih * 0.12), _ih - _nh)
    bg_arr = np.array(_rgb.crop((0, _y0, _iw, _y0 + _nh)))

with PILImage.open(LOGO_PATH) as _logo:
    logo_arr = np.array(_logo.convert("RGBA")).astype(np.float32) / 255.0

H_GRAD = 600; grad = np.zeros((H_GRAD, 1, 4), dtype=np.float32)
for _r in range(H_GRAD):
    _t = _r / (H_GRAD - 1)
    _stops = [(0.00, (0,33,71)), (0.54, (0,70,128)), (1.00, (144,171,196))]
    for _i in range(len(_stops) - 1):
        _t0, _c0 = _stops[_i]; _t1, _c1 = _stops[_i + 1]
        if _t0 <= _t <= _t1:
            _s = (_t - _t0) / (_t1 - _t0)
            grad[_r, 0] = [(_c0[0]+_s*(_c1[0]-_c0[0]))/255,
                           (_c0[1]+_s*(_c1[1]-_c0[1]))/255,
                           (_c0[2]+_s*(_c1[2]-_c0[2]))/255, 0.82]; break

# ── Constantes layout ──────────────────────────────────────────────────────────
SW, SH   = 13.333, 7.5
SW_EMU   = 12192000
SH_EMU   = 6858000

PIC_L, PIC_T, PIC_W, PIC_H = 786581, 1125000, 10599174, 3720000
TITLE_L, TITLE_T, TITLE_W, TITLE_H = PIC_L, 185000, PIC_W, 710000
POP_L,   POP_T,   POP_W,   POP_H   = PIC_L, 845000, 9000000, 255000
LOGO_L,  LOGO_T,  LOGO_W,  LOGO_H  = 9813773, 656354, 1756626, 697725

def _ex(e): return e / SW_EMU
def _ey(e): return e / SH_EMU
def _fig_rect(l, t, w, h): return (l, 1-t-h, w, h)

PIC_RECT  = _fig_rect(_ex(PIC_L), _ey(PIC_T), _ex(PIC_W), _ey(PIC_H))
LOGO_RECT = _fig_rect(_ex(LOGO_L), _ey(LOGO_T), _ex(LOGO_W), _ey(LOGO_H))

# ── Colores ────────────────────────────────────────────────────────────────────
COL_FORM = "#52C97A"   # verde formados
COL_CTRL = "#FFB74D"   # naranja control
COL_MEAN_F = "#FFFFFF" # media formados: blanco
COL_MEAN_C = "#FFD580" # media control: amarillo
DARK_BG  = "#0A0F18"
PERIODOS = ["2023-01","2023-02","2024-01","2024-02","2025-01","2025-02"]
ANIOS_EDD = [2022, 2023, 2024, 2025]

# ── Matplotlib helpers ─────────────────────────────────────────────────────────
def _bg_fig():
    fig = plt.figure(figsize=(SW, SH), facecolor="#101820")
    for z, arr in [(0, bg_arr), (1, grad)]:
        ax = fig.add_axes([0,0,1,1], zorder=z)
        ax.imshow(arr, extent=[0,1,0,1], aspect="auto", origin="upper")
        ax.set_xlim(0,1); ax.set_ylim(0,1); ax.axis("off")
    al = fig.add_axes([LOGO_RECT[0],LOGO_RECT[1],LOGO_RECT[2],LOGO_RECT[3]], zorder=10, facecolor="none")
    al.imshow(logo_arr, aspect="auto"); al.axis("off"); al.patch.set_visible(False)
    return fig

def _tr_fig():
    fig = plt.figure(figsize=(SW, SH), facecolor="none")
    fig.patch.set_facecolor("none")
    return fig

SHARED_BG = os.path.join(OUT_DIR, "_background.png")

def _ensure_bg():
    if not os.path.exists(SHARED_BG):
        fig = _bg_fig()
        plt.savefig(SHARED_BG, dpi=150, facecolor=fig.get_facecolor())
        plt.close()

def _save_ch(fig, name):
    path = os.path.join(OUT_DIR, name)
    plt.savefig(path, dpi=150, facecolor="none", transparent=True)
    plt.close()
    return path

# ── python-pptx helpers ────────────────────────────────────────────────────────
def _new_sl(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def _pic(sl, path, prs):
    sl.shapes.add_picture(path, Emu(0), Emu(0), prs.slide_width, prs.slide_height)

def _txt(sl, text, left, top, width, height, fs=12, bold=False, italic=False,
         color="#FFFFFF", align=PP_ALIGN.LEFT, font_name=None):
    txb = sl.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = txb.text_frame; tf.word_wrap = True
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run(); run.text = line
        run.font.size = Pt(fs); run.font.bold = bold; run.font.italic = italic
        if font_name: run.font.name = font_name
        r, g, b = int(color[1:3],16), int(color[3:5],16), int(color[5:7],16)
        run.font.color.rgb = RGBColor(r,g,b)

def _T(sl, text, fs=20):
    _txt(sl, text, TITLE_L, TITLE_T, TITLE_W, TITLE_H, fs=fs, bold=True, align=PP_ALIGN.CENTER)

def _POP(sl, text):
    _txt(sl, text, POP_L, POP_T, POP_W, POP_H, fs=7.5, italic=True, color="#C8DCF0")

def _CT(sl, text):
    _txt(sl, text, PIC_L, int((_ey(PIC_T) + 0.01)*SH_EMU), PIC_W, int(0.04*SH_EMU),
         fs=9, color="#FFFFFF", font_name="Calibri")

# ── Spaghetti chart ────────────────────────────────────────────────────────────
def _spaghetti(ax, pivot, x_labels, color_ind, color_mean,
               ylabel="", title="", y_fmt=".2f", y_scale=1.0):
    """
    pivot: DataFrame (index=rut, columns=x_labels) — valores ya escalados
    Dibuja líneas individuales tenues + línea de media resaltada.
    """
    x = np.arange(len(x_labels))
    n_plotted = 0

    for _, row in pivot.iterrows():
        vals = row.values.astype(float) * y_scale
        # Solo conecta puntos no-NaN
        mask = ~np.isnan(vals)
        if mask.sum() < 2:
            continue
        ax.plot(x[mask], vals[mask],
                color=color_ind, alpha=0.13, linewidth=0.65, zorder=2)
        n_plotted += 1

    # Media
    means = (pivot.mean(axis=0).values * y_scale)
    stds  = (pivot.std(axis=0).values  * y_scale)

    # Banda ±1σ
    ax.fill_between(x, means - stds, means + stds,
                    color=color_mean, alpha=0.10, zorder=3)

    ax.plot(x, means,
            color=color_mean, linewidth=2.5, marker="o", markersize=7,
            zorder=6, label=f"Promedio (nº {n_plotted})")

    # Etiquetas en la media
    y_range = np.nanmax(means) - np.nanmin(means)
    offset  = max(y_range * 0.04, 0.02)
    for xi, m in enumerate(means):
        if not np.isnan(m):
            ax.text(xi, m + offset, f"{m:{y_fmt}}",
                    ha="center", va="bottom", fontsize=8, fontweight="bold",
                    color=color_mean,
                    path_effects=[pe.withStroke(linewidth=2, foreground=DARK_BG)])

    # Estilo
    ax.set_xticks(x)
    ax.set_xticklabels(x_labels, fontsize=9, color="white", rotation=15, ha="right")
    ax.tick_params(axis="y", colors="#AAAAAA", labelsize=8.5, length=0)
    ax.tick_params(axis="x", length=0)
    for sp in ax.spines.values():
        sp.set_edgecolor("white"); sp.set_alpha(0.22); sp.set_linewidth(0.7)
    ax.yaxis.grid(True, color="white", alpha=0.07, linewidth=0.5)
    ax.set_axisbelow(True)
    ax.set_facecolor("none")
    if ylabel:
        ax.set_ylabel(ylabel, color="#AAAAAA", fontsize=9)
    if title:
        ax.set_title(title, color="white", fontsize=10.5, fontweight="bold", pad=8)
    ax.legend(fontsize=8.5, framealpha=0.22, labelcolor="white",
              facecolor="#101820", edgecolor="#444", loc="upper left")

    return n_plotted

# ── Cargar datos ───────────────────────────────────────────────────────────────
print("Cargando datos…")

sat_210 = pd.read_csv(os.path.join(CASCADE, "05_aptos_p3", "p3_sat_zscore.csv"), encoding="utf-8-sig")
sat_210["rut_key"] = sat_210["rut_key"].astype(str).str.strip()
aptos_ruts = set(sat_210["rut_key"])

scatter = pd.read_csv(os.path.join(COMP, "scatter_sat_notas.csv"), encoding="utf-8-sig")
scatter["rut_docente"] = scatter["rut_docente"].astype(str).str.strip()
scatter["formado"] = scatter["formado"].astype(str).str.upper().isin(["TRUE","1","SI","SÍ","YES"])

edd = pd.read_csv(os.path.join(COMP, "evaluacion_jefes.csv"), encoding="utf-8-sig")
edd["rut_key"] = edd["rut_key"].astype(str).str.strip()
edd["edd_pct"] = pd.to_numeric(edd["edd_total"], errors="coerce") * 100

form_ruts = set(scatter[scatter["formado"]]["rut_docente"].unique())
ctrl_ruts = set(scatter[~scatter["formado"]]["rut_docente"].unique()) - form_ruts

print(f"  Aptos P3:       {len(aptos_ruts)}")
print(f"  Formados (scat):{len(form_ruts)}")
print(f"  Control (scat): {len(ctrl_ruts)}")
print(f"  EDD filas:      {len(edd)}")

# ── Pivots ─────────────────────────────────────────────────────────────────────
def _pivot(df, rut_col, period_col, val_col, filter_ruts=None, min_periods=2):
    df = df.copy()
    if filter_ruts is not None:
        df = df[df[rut_col].isin(filter_ruts)]
    piv = (df.groupby([rut_col, period_col])[val_col]
             .mean()
             .unstack(period_col))
    return piv[piv.count(axis=1) >= min_periods]

# B2 — SAT individual de los 210 aptos por período
piv_b2 = _pivot(scatter, "rut_docente", "periodo", "sat",
                filter_ruts=aptos_ruts).reindex(columns=PERIODOS)
print(f"  B2 SAT pivotado: {len(piv_b2)} docentes × {piv_b2.count().to_dict()}")

# B3 — Nota promedio por período (formados vs control)
piv_b3_f = _pivot(scatter, "rut_docente", "periodo", "nota_promedio",
                  filter_ruts=form_ruts).reindex(columns=PERIODOS)
piv_b3_c = _pivot(scatter, "rut_docente", "periodo", "nota_promedio",
                  filter_ruts=ctrl_ruts).reindex(columns=PERIODOS)
print(f"  B3 Notas form:  {len(piv_b3_f)} docentes")
print(f"  B3 Notas ctrl:  {len(piv_b3_c)} docentes")

# B4 — EDD por año (formados vs control)
piv_b4_f = _pivot(edd, "rut_key", "anio_evaluacion", "edd_pct",
                  filter_ruts=form_ruts).reindex(columns=ANIOS_EDD)
piv_b4_c = _pivot(edd, "rut_key", "anio_evaluacion", "edd_pct",
                  filter_ruts=ctrl_ruts).reindex(columns=ANIOS_EDD)
print(f"  B4 EDD form:    {len(piv_b4_f)} docentes")
print(f"  B4 EDD ctrl:    {len(piv_b4_c)} docentes")

# ── Slides ─────────────────────────────────────────────────────────────────────
def slide_b2(prs):
    """SAT individual: 210 aptos a través de los 6 períodos."""
    fig = _tr_fig()
    L, BY, W, H = PIC_RECT

    ax = fig.add_axes([L + 0.06, BY + 0.05, W - 0.10, H - 0.08], facecolor="none", zorder=5)
    n = _spaghetti(ax, piv_b2, PERIODOS,
                   color_ind=COL_FORM, color_mean=COL_MEAN_F,
                   ylabel="SAT (nota)", title="Evolución SAT individual — 210 Aptos P3",
                   y_fmt=".2f")

    _ensure_bg()
    sl = _new_sl(prs); _pic(sl, SHARED_BG, prs); _pic(sl, _save_ch(fig, "tray_b2_sat.png"), prs)
    _T(sl, "Bloque II — Evolución SAT Individual por Período (Aptos P3)")
    _POP(sl, f"Cada línea = 1 docente (nº {n} con ≥2 períodos)  ·  Períodos: 2023-01 a 2025-02  ·  Banda sombreada = ±1 DE")
    _CT(sl, "SAT ponderado por nº alumnos evaluaron (cobertura ≥40%)  ·  Línea blanca = promedio del grupo")
    print("  ✓ Slide B2 — SAT individual")


def slide_b3(prs):
    """Nota promedio alumnos: formados vs control — dos paneles pareados."""
    fig = _tr_fig()
    L, BY, W, H = PIC_RECT
    half_w = W / 2 - 0.03

    ax_f = fig.add_axes([L + 0.04, BY + 0.05, half_w, H - 0.08], facecolor="none", zorder=5)
    ax_c = fig.add_axes([L + half_w + 0.06, BY + 0.05, half_w, H - 0.08], facecolor="none", zorder=5)

    # Calcular rango compartido para ejes Y comparables
    all_vals = pd.concat([piv_b3_f, piv_b3_c]).values.flatten()
    y_min = np.nanpercentile(all_vals, 2) - 0.05
    y_max = np.nanpercentile(all_vals, 98) + 0.05

    nf = _spaghetti(ax_f, piv_b3_f, PERIODOS,
                    color_ind=COL_FORM, color_mean=COL_MEAN_F,
                    ylabel="Nota promedio alumnos (1–7)",
                    title=f"Formados (nº {len(piv_b3_f)})", y_fmt=".2f")
    nc = _spaghetti(ax_c, piv_b3_c, PERIODOS,
                    color_ind=COL_CTRL, color_mean=COL_MEAN_C,
                    title=f"Control (nº {len(piv_b3_c)})", y_fmt=".2f")

    # Mismo rango Y en ambos paneles
    ax_f.set_ylim(y_min, y_max)
    ax_c.set_ylim(y_min, y_max)
    ax_c.tick_params(axis="y", left=False, labelleft=False)

    # Divisor vertical
    div = fig.add_axes([L + half_w + 0.025, BY + 0.02, 0.0015, H - 0.04],
                       facecolor="white", alpha=0.18, zorder=3)
    div.axis("off")

    _ensure_bg()
    sl = _new_sl(prs); _pic(sl, SHARED_BG, prs); _pic(sl, _save_ch(fig, "tray_b3_notas.png"), prs)
    _T(sl, "Bloque III — Evolución Notas Alumnos por Período")
    _POP(sl, f"Cada línea = 1 docente  ·  Izq.: formados (nº {nf})  ·  Der.: control (nº {nc})  ·  Nota media de secciones del docente por período")
    _CT(sl, "Nota promedio de alumnos aggregada por docente × período  ·  Escala común en ambos paneles  ·  Banda = ±1 DE")
    print("  ✓ Slide B3 — Notas alumnos pareado")


def slide_b4(prs):
    """EDD jefaturas: formados vs control — dos paneles pareados."""
    fig = _tr_fig()
    L, BY, W, H = PIC_RECT
    half_w = W / 2 - 0.03

    ax_f = fig.add_axes([L + 0.04, BY + 0.05, half_w, H - 0.08], facecolor="none", zorder=5)
    ax_c = fig.add_axes([L + half_w + 0.06, BY + 0.05, half_w, H - 0.08], facecolor="none", zorder=5)

    all_vals = pd.concat([piv_b4_f, piv_b4_c]).values.flatten()
    y_min = max(0, np.nanpercentile(all_vals, 2) - 3)
    y_max = min(100, np.nanpercentile(all_vals, 98) + 3)

    x_lbls = [str(a) for a in ANIOS_EDD]

    nf = _spaghetti(ax_f, piv_b4_f, x_lbls,
                    color_ind=COL_FORM, color_mean=COL_MEAN_F,
                    ylabel="EDD total (%)",
                    title=f"Formados (nº {len(piv_b4_f)})", y_fmt=".0f")
    nc = _spaghetti(ax_c, piv_b4_c, x_lbls,
                    color_ind=COL_CTRL, color_mean=COL_MEAN_C,
                    title=f"Control (nº {len(piv_b4_c)})", y_fmt=".0f")

    ax_f.set_ylim(y_min, y_max)
    ax_c.set_ylim(y_min, y_max)
    ax_c.tick_params(axis="y", left=False, labelleft=False)

    div = fig.add_axes([L + half_w + 0.025, BY + 0.02, 0.0015, H - 0.04],
                       facecolor="white", alpha=0.18, zorder=3)
    div.axis("off")

    _ensure_bg()
    sl = _new_sl(prs); _pic(sl, SHARED_BG, prs); _pic(sl, _save_ch(fig, "tray_b4_edd.png"), prs)
    _T(sl, "Bloque IV — Evolución EDD Jefaturas por Año")
    _POP(sl, f"Cada línea = 1 docente  ·  Izq.: formados (nº {nf})  ·  Der.: control (nº {nc})  ·  EDD total expresado en %")
    _CT(sl, "Evaluación jefaturas: edd_total × 100  ·  Escala común en ambos paneles  ·  Banda = ±1 DE  ·  Años: 2022–2025")
    print("  ✓ Slide B4 — EDD jefaturas pareado")


# ── Main ───────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    print("\nGenerando TRAYECTORIAS_v1.pptx (3 slides)…")
    _ensure_bg()
    prs = Presentation()
    prs.slide_width  = Emu(SW_EMU)
    prs.slide_height = Emu(SH_EMU)

    slide_b2(prs)
    slide_b3(prs)
    slide_b4(prs)

    prs.save(OUT_PPTX)
    print(f"\n✓ Guardado: {OUT_PPTX}")
    print(f"  {len(prs.slides)} diapositivas")
