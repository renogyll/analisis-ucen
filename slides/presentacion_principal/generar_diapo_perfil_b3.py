"""
Genera UNA diapo: Bloque III — Perfil Demográfico Formados vs Control
Salida: outputs/pptx/DIAPO_perfil_b3.pptx
"""
import sys; sys.stdout.reconfigure(encoding="utf-8")
import os, zipfile
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.patheffects as pe
import pandas as pd
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import pathlib

# ── Rutas ─────────────────────────────────────────────────────────────────────
BASE     = os.path.dirname(os.path.abspath(__file__))
REPO     = str(pathlib.Path(BASE).parents[1])
CASCADE  = os.path.join(REPO, "data", "cascade")
COMP     = os.path.join(CASCADE, "complementarios")

SAT_CSV   = os.path.join(CASCADE, "05_aptos_p3", "p3_sat_zscore.csv")
SCAT_CSV  = os.path.join(COMP, "scatter_sat_notas.csv")
NOM_CSV   = os.path.join(CASCADE, "00_base", "nomina_x_dotacion.csv")
FONDOTIPO = (r"c:\Users\r.gonzalez_fluxsolar.LAPTOP-FLUX-ECO"
             r"\Downloads\Analisis_UCEN_v2\Fondotipop.pptx")
OUT_PPTX  = os.path.join(REPO, "outputs", "pptx", "DIAPO_perfil_b3.pptx")
OUT_DIR   = os.path.join(BASE, "dark_slides_v3")
SCRATCH   = os.path.join(REPO, "outputs", "scratch")
os.makedirs(OUT_DIR, exist_ok=True)
os.makedirs(SCRATCH, exist_ok=True)

# ── Layout EMU ────────────────────────────────────────────────────────────────
SW, SH       = 13.333, 7.5
SW_EMU       = 12192000
SH_EMU       = 6858000
PIC_L, PIC_T, PIC_W, PIC_H = 786581, 1125000, 10599174, 3720000
BUL_L, BUL_T, BUL_W, BUL_H = 786581, 4870000, 10599174, 1870000
LOGO_L, LOGO_T, LOGO_W, LOGO_H = 9813773, 656354, 1756626, 697725
TITLE_L, TITLE_T, TITLE_W, TITLE_H = PIC_L, 185000, PIC_W, 710000
POP_L,   POP_T,   POP_W,   POP_H   = PIC_L, 845000, 9000000, 255000

def _ex(e): return e / SW_EMU
def _ey(e): return e / SH_EMU
def _fig_rect(l, t, w, h): return (l, 1 - t - h, w, h)

PIC_RECT  = _fig_rect(_ex(PIC_L), _ey(PIC_T), _ex(PIC_W), _ey(PIC_H))
LOGO_RECT = _fig_rect(_ex(LOGO_L), _ey(LOGO_T), _ex(LOGO_W), _ey(LOGO_H))

# ── Colores ───────────────────────────────────────────────────────────────────
COL_FORM_P = "#5C9BD6"
COL_CTRL_P = "#FFB74D"

# ── Categorías ────────────────────────────────────────────────────────────────
JER_ORD = ["INSTRUCTOR REGULAR", "INSTRUCTOR DOCENTE",
           "ASISTENTE REGULAR",  "ASISTENTE DOCENTE",
           "ASOCIADO REGULAR",   "ASOCIADO DOCENTE",
           "TITULAR REGULAR",    "TITULAR DOCENTE"]
JER_LBL = ["Instr. Regular", "Instr. Docente",
            "Asist. Regular", "Asist. Docente",
            "Asoc. Regular",  "Asoc. Docente",
            "Tit. Regular",   "Tit. Docente"]
TRAMOS_EDAD  = ["< 35", "35–44", "45–54", "55–64", "≥ 65"]
TRAMOS_EDAD_RAW = {
    "<30": "< 35",  "30-34": "< 35",
    "35-39": "35–44", "40-44": "35–44",
    "45-49": "45–54", "50-54": "45–54",
    "55-59": "55–64", "60-64": "55–64",
    "65-69": "≥ 65",  "70+":   "≥ 65",
}

# ── Assets ────────────────────────────────────────────────────────────────────
BG_PATH   = os.path.join(SCRATCH, "fondotipo_image1.jpg")
LOGO_PATH = os.path.join(SCRATCH, "fondotipo_image2.png")
for path, zname in [(BG_PATH, "ppt/media/image1.jpg"), (LOGO_PATH, "ppt/media/image2.png")]:
    if not os.path.exists(path):
        with zipfile.ZipFile(FONDOTIPO) as z:
            with open(path, "wb") as f:
                f.write(z.read(zname))

with PILImage.open(BG_PATH) as _im:
    _rgb = _im.convert("RGB"); _iw, _ih = _rgb.size
    _nh  = int(_iw / (16 / 9)); _y0 = min(int(_ih * 0.12), _ih - _nh)
    bg_arr = np.array(_rgb.crop((0, _y0, _iw, _y0 + _nh)))

with PILImage.open(LOGO_PATH) as _logo:
    logo_arr = np.array(_logo.convert("RGBA")).astype(np.float32) / 255.0

H_GRAD = 600; grad = np.zeros((H_GRAD, 1, 4), dtype=np.float32)
for _r in range(H_GRAD):
    _t = _r / (H_GRAD - 1)
    _stops = [(0.00, (0, 33, 71)), (0.54, (0, 70, 128)), (1.00, (144, 171, 196))]
    for _i in range(len(_stops) - 1):
        _t0, _c0 = _stops[_i]; _t1, _c1 = _stops[_i + 1]
        if _t0 <= _t <= _t1:
            _s = (_t - _t0) / (_t1 - _t0)
            grad[_r, 0] = [(_c0[0] + _s * (_c1[0] - _c0[0])) / 255,
                           (_c0[1] + _s * (_c1[1] - _c0[1])) / 255,
                           (_c0[2] + _s * (_c1[2] - _c0[2])) / 255, 0.82]; break

# ── Helpers figura ─────────────────────────────────────────────────────────────
SHARED_BG = os.path.join(OUT_DIR, "_background.png")

def _bg_fig():
    fig = plt.figure(figsize=(SW, SH), facecolor="#101820")
    fig.patch.set_facecolor("#101820")
    for z, arr in [(0, bg_arr), (1, grad)]:
        ax = fig.add_axes([0, 0, 1, 1], zorder=z)
        ax.imshow(arr, extent=[0, 1, 0, 1], aspect="auto", origin="upper")
        ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
    al = fig.add_axes([LOGO_RECT[0], LOGO_RECT[1], LOGO_RECT[2], LOGO_RECT[3]],
                      zorder=10, facecolor="none")
    al.imshow(logo_arr, aspect="auto"); al.axis("off"); al.patch.set_visible(False)
    return fig

def _tr_fig():
    fig = plt.figure(figsize=(SW, SH), facecolor="none")
    fig.patch.set_facecolor("none")
    return fig

def _save_bg(fig, name):
    path = os.path.join(OUT_DIR, name)
    fig.savefig(path, dpi=150, facecolor=fig.get_facecolor())
    plt.close(fig); return path

def _save_ch(fig, name):
    path = os.path.join(OUT_DIR, name)
    fig.savefig(path, dpi=150, facecolor="none", transparent=True)
    plt.close(fig); return path

def _ensure_bg():
    if not os.path.exists(SHARED_BG):
        _save_bg(_bg_fig(), "_background.png")

# ── Helpers pptx ──────────────────────────────────────────────────────────────
def _new_sl(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def _pic(sl, path, prs):
    sl.shapes.add_picture(path, Emu(0), Emu(0), prs.slide_width, prs.slide_height)

def _txt(sl, text, left, top, width, height,
         fs=12, bold=False, italic=False, color="#FFFFFF",
         align=PP_ALIGN.LEFT, wrap=True, font_name=None):
    txb = sl.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf  = txb.text_frame; tf.word_wrap = wrap
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run(); run.text = line
        run.font.size = Pt(fs); run.font.bold = bold; run.font.italic = italic
        if font_name: run.font.name = font_name
        r, g, b = int(color[1:3], 16), int(color[3:5], 16), int(color[5:7], 16)
        run.font.color.rgb = RGBColor(r, g, b)

def _T(sl, text, fs=20):
    _txt(sl, text, TITLE_L, TITLE_T, TITLE_W, TITLE_H,
         fs=fs, bold=True, color="#FFFFFF", align=PP_ALIGN.CENTER)

def _POP(sl, text):
    _txt(sl, text, POP_L, POP_T, POP_W, POP_H, fs=7.5, italic=True, color="#C8DCF0")

def _BUL(sl, items, fs=11.5):
    txb = sl.shapes.add_textbox(Emu(BUL_L), Emu(BUL_T), Emu(BUL_W), Emu(BUL_H))
    tf  = txb.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(5); p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = f"{i + 1}.  {item}"
        run.font.size = Pt(fs)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

def _style_ax(ax):
    ax.tick_params(axis="x", colors="#AAAAAA", labelsize=8, length=0)
    ax.tick_params(axis="y", colors="white",   labelsize=8, length=0)
    for sp in ax.spines.values():
        sp.set_edgecolor("white"); sp.set_alpha(0.20); sp.set_linewidth(0.7)
    ax.set_axisbelow(True)

# ── Cálculo de porcentajes (denominador = solo docentes con dato reconocido) ──
def _pct_jer(df, col="jerarquia_u"):
    vc = df[col].value_counts()
    n  = max(int(vc[vc.index.isin(JER_ORD)].sum()), 1)
    return [round(vc.get(j, 0) / n * 100, 1) for j in JER_ORD]

def _pct_tramo(df, col="tramo_g"):
    vc = df[col].value_counts()
    n  = max(int(vc.sum()), 1)
    return [round(vc.get(t, 0) / n * 100, 1) for t in TRAMOS_EDAD]

# ── Butterfly chart ────────────────────────────────────────────────────────────
def _butterfly(fig, bx, by, bw, bh, cats, lbl, pct_f, pct_c,
               n_f_con_dato, n_c_con_dato, n_f_tot, n_c_tot, chart_title):
    ya      = np.arange(len(cats))
    max_val = max(max(pct_f, default=1), max(pct_c, default=1))
    xlim    = max_val * 1.35
    gap_lbl = 0.045
    w_panel = (bw - gap_lbl) / 2

    # Panel izquierdo — Formados
    ax_f = fig.add_axes([bx, by, w_panel, bh], facecolor="none", zorder=5)
    ax_f.barh(ya, pct_f, height=0.60, color=COL_FORM_P, alpha=0.88, edgecolor="none")
    ax_f.set_xlim(xlim, 0)
    ax_f.set_ylim(-0.6, len(cats) - 0.4)
    ax_f.set_yticks(ya); ax_f.set_yticklabels([])
    ax_f.xaxis.grid(True, color="white", alpha=0.06, linewidth=0.5)
    for j, v in enumerate(pct_f):
        if v >= 4:
            ax_f.text(v / 2, j, f"{v:.0f}%", ha="center", va="center",
                      fontsize=7.5, color="white", fontweight="bold")
        elif v > 0:
            ax_f.text(v + 0.5, j, f"{v:.0f}%", ha="right", va="center",
                      fontsize=7, color="#AAAAAA")
    ax_f.xaxis.set_major_formatter(plt.FuncFormatter(lambda x, _: f"{abs(x):.0f}%"))
    ax_f.set_title(f"Formados  (nº {n_f_con_dato})", color=COL_FORM_P,
                   fontsize=9, fontweight="bold", pad=5)
    _style_ax(ax_f)

    # Panel derecho — Control
    ax_c = fig.add_axes([bx + w_panel + gap_lbl, by, w_panel, bh], facecolor="none", zorder=5)
    ax_c.barh(ya, pct_c, height=0.60, color=COL_CTRL_P, alpha=0.88, edgecolor="none")
    ax_c.set_xlim(0, xlim)
    ax_c.set_ylim(-0.6, len(cats) - 0.4)
    ax_c.set_yticks(ya); ax_c.set_yticklabels(lbl, fontsize=8.5, color="white")
    ax_c.yaxis.set_ticks_position("left")
    ax_c.xaxis.grid(True, color="white", alpha=0.06, linewidth=0.5)
    for j, v in enumerate(pct_c):
        if v >= 4:
            ax_c.text(v / 2, j, f"{v:.0f}%", ha="center", va="center",
                      fontsize=7.5, color="white", fontweight="bold")
        elif v > 0:
            ax_c.text(v + 0.5, j, f"{v:.0f}%", ha="left", va="center",
                      fontsize=7, color="#AAAAAA")
    ax_c.xaxis.set_major_formatter(plt.FuncFormatter(lambda x, _: f"{x:.0f}%"))
    ax_c.set_title(f"Control  (nº {n_c_con_dato})", color=COL_CTRL_P,
                   fontsize=9, fontweight="bold", pad=5)
    _style_ax(ax_c)

    fig.text(bx + bw / 2, by + bh + 0.035, chart_title,
             ha="center", va="bottom", fontsize=10, color="white", fontweight="bold",
             path_effects=[pe.withStroke(linewidth=1.5, foreground="#0A1830")])

# ── Carga de datos ─────────────────────────────────────────────────────────────
sat = pd.read_csv(SAT_CSV, encoding="utf-8-sig")
sat["rut_key"]     = sat["rut_key"].astype(str).str.strip()
sat["jerarquia_u"] = sat["jerarquia"].str.strip().str.upper()
sat["tramo_g"]     = sat["tramo_edad"].map(TRAMOS_EDAD_RAW)
N_FORM = len(sat)

scat = pd.read_csv(SCAT_CSV, encoding="utf-8-sig")
scat["formado"] = (scat["formado"].astype(str).str.strip().str.upper()
                   .isin(["TRUE", "1", "SI", "SÍ", "YES"]))
scat_ctrl_ruts = set(scat[~scat["formado"]]["rut_docente"].astype(str).str.strip().unique())

nom = pd.read_csv(NOM_CSV, encoding="utf-8-sig")
nom["rut_key"] = nom["rut_key"].astype(str).str.strip()
ctrl = nom[nom["rut_key"].isin(scat_ctrl_ruts)].drop_duplicates("rut_key").copy()
ctrl["jerarquia_u"] = ctrl["jerarquia"].str.strip().str.upper()
ctrl["tramo_g"]     = ctrl["tramo_edad"].map(TRAMOS_EDAD_RAW)
N_CTRL = len(ctrl)

# n con dato reconocido por categoría
n_f_jer = int(sat["jerarquia_u"].isin(JER_ORD).sum())
n_c_jer = int(ctrl["jerarquia_u"].isin(JER_ORD).sum())
n_f_tra = int(sat["tramo_g"].notna().sum())
n_c_tra = int(ctrl["tramo_g"].notna().sum())

pct_f_jer = _pct_jer(sat)
pct_c_jer = _pct_jer(ctrl)
pct_f_tra = _pct_tramo(sat)
pct_c_tra = _pct_tramo(ctrl)

print(f"Formados:  {N_FORM}  (jer con dato: {n_f_jer}, tramo con dato: {n_f_tra})")
print(f"Control:   {N_CTRL}  (jer con dato: {n_c_jer}, tramo con dato: {n_c_tra})")
print(f"Jerarquía sum Formados={sum(pct_f_jer):.1f}%  Control={sum(pct_c_jer):.1f}%")
print(f"Tramo     sum Formados={sum(pct_f_tra):.1f}%  Control={sum(pct_c_tra):.1f}%")

# ── Generar gráfico ─────────────────────────────────────────────────────────────
_ensure_bg()

px = PIC_RECT[0] + 0.01;  py = PIC_RECT[1] + 0.02
pw = PIC_RECT[2] - 0.02;  ph = 0.42
GAP = 0.025;  bw = (pw - GAP) / 2

fig = _tr_fig()

_butterfly(fig, bx=px,          by=py, bw=bw, bh=ph,
           cats=JER_ORD, lbl=JER_LBL,
           pct_f=pct_f_jer, pct_c=pct_c_jer,
           n_f_con_dato=n_f_jer, n_c_con_dato=n_c_jer,
           n_f_tot=N_FORM,       n_c_tot=N_CTRL,
           chart_title="Jerarquía Académica")

_butterfly(fig, bx=px + bw + GAP, by=py, bw=bw, bh=ph,
           cats=TRAMOS_EDAD, lbl=TRAMOS_EDAD,
           pct_f=pct_f_tra, pct_c=pct_c_tra,
           n_f_con_dato=n_f_tra, n_c_con_dato=n_c_tra,
           n_f_tot=N_FORM,        n_c_tot=N_CTRL,
           chart_title="Tramo de Edad")

chart_path = _save_ch(fig, "perfil_b3.png")

# ── Bullet dinámico ────────────────────────────────────────────────────────────
jer_dif = sorted(
    [(JER_LBL[i], pct_f_jer[i] - pct_c_jer[i]) for i in range(len(JER_ORD))],
    key=lambda x: -abs(x[1]))
eda_dif = sorted(
    [(TRAMOS_EDAD[i], pct_f_tra[i] - pct_c_tra[i]) for i in range(len(TRAMOS_EDAD))],
    key=lambda x: -abs(x[1]))

# ── Ensamblar PPTX ─────────────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width  = Emu(SW_EMU)
prs.slide_height = Emu(SH_EMU)

sl = _new_sl(prs)
_pic(sl, SHARED_BG, prs)
_pic(sl, chart_path, prs)

_T(sl, "Bloque III — Perfil Demográfico: Formados vs Control  (Rendimiento Académico)", fs=17)
_POP(sl, f"Formados Aptos P3: nº {N_FORM}  ·  Grupo Control sin formación: nº {N_CTRL}")
_BUL(sl, [
    f"Jerarquía: '{jer_dif[0][0]}' tiene la mayor diferencia entre grupos ({jer_dif[0][1]:+.0f} pp). "
    "Las barras muestran % dentro de cada grupo — no volumen absoluto — para hacer comparables "
    "grupos de distinto tamaño.",
    f"Edad: '{eda_dif[0][0]}' es el tramo más diferenciado ({eda_dif[0][1]:+.0f} pp). "
    "Si los grupos difieren en composición, parte de la brecha de aprobación puede deberse "
    "a eso y no solo a la formación.",
])

prs.save(OUT_PPTX)
print(f"\n✓ Guardado: {OUT_PPTX}")
