"""
generar_poblaciones.py
3 slides: distribuciones de Tipo Contrato, Jerarquía y Sexo
para las poblaciones Universo Base (1.144) / Control B2 (568) / Aptos P3 (210).
Ejecutar desde el directorio analisis-ucen/ o directamente.
"""
import sys; sys.stdout.reconfigure(encoding="utf-8")
import os, zipfile, io
from pathlib import Path
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patheffects as pe
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from config import CASCADE

# ── Paths ──────────────────────────────────────────────────────────────────────
REPO      = Path(__file__).resolve().parents[2]
SCRATCH   = os.path.join(REPO, "outputs", "scratch")
OUT_DIR   = os.path.join(REPO, "outputs", "scratch", "pop_slides")
OUT_PPTX  = os.path.join(REPO, "outputs", "pptx", "POBLACIONES_v1.pptx")
FONDOTIPO = os.path.join(REPO, "PRESENTACION_800.pptx")
os.makedirs(OUT_DIR, exist_ok=True)
os.makedirs(os.path.join(REPO, "outputs", "pptx"), exist_ok=True)

COMP = os.path.join(CASCADE, "complementarios")

# ── Assets: fondo + logo (mismo que presentación principal) ───────────────────
BG_PATH   = os.path.join(SCRATCH, "fondotipo_image1.jpg")
LOGO_PATH = os.path.join(SCRATCH, "fondotipo_image2.png")
for path, zname in [(BG_PATH, "ppt/media/image1.jpg"), (LOGO_PATH, "ppt/media/image2.png")]:
    if not os.path.exists(path):
        with zipfile.ZipFile(FONDOTIPO) as z:
            with open(path, "wb") as f:
                f.write(z.read(zname))

with PILImage.open(BG_PATH) as _im:
    _rgb = _im.convert("RGB"); _iw, _ih = _rgb.size
    _nh  = int(_iw / (16/9)); _y0 = min(int(_ih * 0.12), _ih - _nh)
    bg_arr = np.array(_rgb.crop((0, _y0, _iw, _y0 + _nh)))

with PILImage.open(LOGO_PATH) as _logo:
    logo_arr = np.array(_logo.convert("RGBA")).astype(np.float32) / 255.0

H_GRAD = 600; grad = np.zeros((H_GRAD, 1, 4), dtype=np.float32)
for _r in range(H_GRAD):
    _t = _r / (H_GRAD - 1)
    _stops = [(0.00, (0, 33, 71)), (0.54, (0, 70, 128)), (1.00, (144, 171, 196))]
    for _i in range(len(_stops) - 1):
        _t0, _c0 = _stops[_i]; _t1, _c1 = _stops[_i + 1]
        if _t0 <= _t <= _t1:
            _s = (_t - _t0) / (_t1 - _t0)
            grad[_r, 0] = [(_c0[0]+_s*(_c1[0]-_c0[0]))/255,
                           (_c0[1]+_s*(_c1[1]-_c0[1]))/255,
                           (_c0[2]+_s*(_c1[2]-_c0[2]))/255, 0.82]; break

# ── Constantes layout ──────────────────────────────────────────────────────────
SW, SH     = 13.333, 7.5
SW_EMU     = 12192000
SH_EMU     = 6858000

PIC_L, PIC_T, PIC_W, PIC_H = 786581, 1125000, 10599174, 3720000
TITLE_L, TITLE_T, TITLE_W, TITLE_H = PIC_L, 185000, PIC_W, 710000
POP_L, POP_T, POP_W, POP_H = PIC_L, 845000, 9000000, 255000
LOGO_L, LOGO_T, LOGO_W, LOGO_H = 9813773, 656354, 1756626, 697725

def _ex(e): return e / SW_EMU
def _ey(e): return e / SH_EMU
def _fig_rect(l, t, w, h): return (l, 1 - t - h, w, h)

PIC_RECT  = _fig_rect(_ex(PIC_L), _ey(PIC_T), _ex(PIC_W), _ey(PIC_H))
LOGO_RECT = _fig_rect(_ex(LOGO_L), _ey(LOGO_T), _ex(LOGO_W), _ey(LOGO_H))

# ── Colores ────────────────────────────────────────────────────────────────────
COL_J   = "#5C9BD6"; COL_H = "#FFB74D"
COL_M   = "#5C9BD6"; COL_F = "#CE93D8"
DARK_BG = "#0A0F18"
JER_PAL = ["#5C9BD6", "#52C97A", "#FFB74D", "#CE93D8", "#80DEEA", "#E57373", "#A47BD6", "#AAAAAA"]

# ── Matplotlib helpers ─────────────────────────────────────────────────────────
def _bg_fig():
    fig = plt.figure(figsize=(SW, SH), facecolor="#101820")
    for z, arr in [(0, bg_arr), (1, grad)]:
        ax = fig.add_axes([0, 0, 1, 1], zorder=z)
        ax.imshow(arr, extent=[0, 1, 0, 1], aspect="auto", origin="upper")
        ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
    al = fig.add_axes([LOGO_RECT[0], LOGO_RECT[1], LOGO_RECT[2], LOGO_RECT[3]],
                     zorder=10, facecolor="none")
    al.imshow(logo_arr, aspect="auto"); al.axis("off"); al.patch.set_visible(False)
    return fig

def _tr_fig():
    fig = plt.figure(figsize=(SW, SH), facecolor="none")
    fig.patch.set_facecolor("none")
    return fig

SHARED_BG = os.path.join(OUT_DIR, "_background.png")

def _ensure_bg():
    if not os.path.exists(SHARED_BG):
        fig = _bg_fig()
        plt.savefig(SHARED_BG, dpi=150, facecolor=fig.get_facecolor())
        plt.close()

def _save_ch(fig, name):
    path = os.path.join(OUT_DIR, name)
    plt.savefig(path, dpi=150, facecolor="none", transparent=True)
    plt.close()
    return path

# ── python-pptx helpers ────────────────────────────────────────────────────────
def _new_sl(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def _pic(sl, path, prs):
    sl.shapes.add_picture(path, Emu(0), Emu(0), prs.slide_width, prs.slide_height)

def _txt(sl, text, left, top, width, height,
         fs=12, bold=False, italic=False, color="#FFFFFF",
         align=PP_ALIGN.LEFT, wrap=True, font_name=None):
    txb = sl.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf  = txb.text_frame; tf.word_wrap = wrap
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run(); run.text = line
        run.font.size = Pt(fs); run.font.bold = bold; run.font.italic = italic
        if font_name: run.font.name = font_name
        r, g, b = int(color[1:3], 16), int(color[3:5], 16), int(color[5:7], 16)
        run.font.color.rgb = RGBColor(r, g, b)

def _T(sl, text, fs=20):
    _txt(sl, text, TITLE_L, TITLE_T, TITLE_W, TITLE_H,
         fs=fs, bold=True, align=PP_ALIGN.CENTER)

def _POP(sl, text):
    _txt(sl, text, POP_L, POP_T, POP_W, POP_H,
         fs=7.5, italic=True, color="#C8DCF0")

def _CT(sl, text):
    ct_l = int(_ex(PIC_L + 500000) * SW_EMU)
    ct_t = int((_ey(PIC_T) + 0.01) * SH_EMU)
    ct_w = int(_ex(PIC_W - 500000) * SW_EMU)
    _txt(sl, text, ct_l, ct_t, ct_w, int(0.04 * SH_EMU),
         fs=9, color="#FFFFFF", font_name="Calibri")

# ── Cargar datos ───────────────────────────────────────────────────────────────
print("Cargando datos…")
base = pd.read_csv(os.path.join(CASCADE, "00_base", "nomina_x_dotacion.csv"), encoding="utf-8-sig")
base["rut_key"] = base["rut_key"].astype(str).str.strip()

sat = pd.read_csv(os.path.join(CASCADE, "05_aptos_p3", "p3_sat_zscore.csv"), encoding="utf-8-sig")
sat["rut_key"] = sat["rut_key"].astype(str).str.strip()

ctrl_raw = pd.read_csv(os.path.join(COMP, "control_918.csv"), encoding="utf-8-sig")
ctrl_raw["rut_key"] = ctrl_raw["rut_key"].astype(str).str.strip()
ctrl = (ctrl_raw.drop_duplicates("rut_key")
        .merge(base[["rut_key", "tipo_contrato_tag", "jerarquia", "sexo"]],
               on="rut_key", how="left"))

print(f"  Universo base: {len(base):,}")
print(f"  Control B2:    {len(ctrl):,}")
print(f"  Aptos P3:      {len(sat):,}")

# ── Funciones de tarta ─────────────────────────────────────────────────────────
def _pie_contrato(ax, df):
    c = df["tipo_contrato_tag"].str.upper().value_counts() if "tipo_contrato_tag" in df.columns else pd.Series(dtype=int)
    nj = int(c.get("JORNADA", 0)); nh = int(c.get("HONORARIO", 0))
    tot = nj + nh
    if tot == 0:
        ax.text(0, 0, "Sin datos", ha="center", va="center", color="#AAA"); ax.axis("off"); return
    pj = nj / tot * 100; ph = nh / tot * 100
    wedges, _ = ax.pie([pj, ph], colors=[COL_J, COL_H],
                       startangle=90, counterclock=False,
                       wedgeprops=dict(linewidth=1.5, edgecolor=DARK_BG, alpha=0.92))
    for w_, val, cnt in [(wedges[0], pj, nj), (wedges[1], ph, nh)]:
        if val >= 5:
            ang = np.radians((w_.theta1 + w_.theta2) / 2)
            x = 0.60 * np.cos(ang); y = 0.60 * np.sin(ang)
            ax.text(x, y, f"{val:.0f}%\n({cnt})",
                    ha="center", va="center", fontsize=9.5, fontweight="bold", color="white",
                    path_effects=[pe.withStroke(linewidth=2.5, foreground=DARK_BG)])
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(facecolor=COL_J, alpha=0.9, label="Jornada"),
                       Patch(facecolor=COL_H, alpha=0.9, label="Honorario")],
              fontsize=8.5, loc="lower center", bbox_to_anchor=(0.5, -0.18),
              ncol=2, framealpha=0.25, labelcolor="white",
              facecolor="#101820", edgecolor="#444")
    ax.set_title("Tipo Contrato", color="white", fontsize=11, fontweight="bold", pad=10)
    ax.set_aspect("equal")


def _pie_jerarquia(ax, df, max_cat=7):
    if "jerarquia" not in df.columns:
        ax.axis("off"); return
    c = df["jerarquia"].str.strip().str.title().value_counts()
    if c.empty:
        ax.axis("off"); return
    top   = c.head(max_cat)
    other = c.iloc[max_cat:].sum()
    if other > 0:
        top = pd.concat([top, pd.Series({"Otros": other})])
    vals = top.values; lbls = top.index.tolist()
    cols = JER_PAL[:len(vals)]
    tot  = vals.sum()
    wedges, _ = ax.pie(vals, colors=cols,
                       startangle=90, counterclock=False,
                       wedgeprops=dict(linewidth=1.2, edgecolor=DARK_BG, alpha=0.90))
    for w_, v in zip(wedges, vals):
        pct = v / tot * 100
        if pct >= 6:
            ang = np.radians((w_.theta1 + w_.theta2) / 2)
            x = 0.60 * np.cos(ang); y = 0.60 * np.sin(ang)
            ax.text(x, y, f"{pct:.0f}%",
                    ha="center", va="center", fontsize=9, fontweight="bold", color="white",
                    path_effects=[pe.withStroke(linewidth=2, foreground=DARK_BG)])
    ax.legend(wedges, [f"{l} ({int(v)})" for l, v in zip(lbls, vals)],
              fontsize=7.5, loc="lower center", bbox_to_anchor=(0.5, -0.32),
              ncol=2, framealpha=0.25, labelcolor="white",
              facecolor="#101820", edgecolor="#444")
    ax.set_title("Jerarquía Académica", color="white", fontsize=11, fontweight="bold", pad=10)
    ax.set_aspect("equal")


def _pie_sexo(ax, df):
    if "sexo" not in df.columns:
        ax.axis("off"); return
    c   = df["sexo"].str.upper().value_counts()
    nm  = int(c.get("HOMBRE", c.get("M", c.get("MASCULINO", 0))))
    nf  = int(c.get("MUJER",  c.get("F", c.get("FEMENINO",  0))))
    tot = nm + nf
    if tot == 0:
        ax.axis("off"); return
    pm = nm / tot * 100; pf = nf / tot * 100
    wedges, _ = ax.pie([pm, pf], colors=[COL_M, COL_F],
                       startangle=90, counterclock=False,
                       wedgeprops=dict(linewidth=1.5, edgecolor=DARK_BG, alpha=0.92))
    for w_, val, cnt, lbl in [(wedges[0], pm, nm, "Hombre"), (wedges[1], pf, nf, "Mujer")]:
        if val >= 5:
            ang = np.radians((w_.theta1 + w_.theta2) / 2)
            x = 0.60 * np.cos(ang); y = 0.60 * np.sin(ang)
            ax.text(x, y, f"{val:.0f}%\n({cnt})",
                    ha="center", va="center", fontsize=9.5, fontweight="bold", color="white",
                    path_effects=[pe.withStroke(linewidth=2.5, foreground=DARK_BG)])
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(facecolor=COL_M, alpha=0.9, label="Hombre"),
                       Patch(facecolor=COL_F, alpha=0.9, label="Mujer")],
              fontsize=8.5, loc="lower center", bbox_to_anchor=(0.5, -0.18),
              ncol=2, framealpha=0.25, labelcolor="white",
              facecolor="#101820", edgecolor="#444")
    ax.set_title("Sexo", color="white", fontsize=11, fontweight="bold", pad=10)
    ax.set_aspect("equal")


# ── Generador de slide ─────────────────────────────────────────────────────────
def make_slide(prs, df, slide_title, pop_text, ct_text, fname):
    fig = _tr_fig()

    L, BY, W, H = PIC_RECT
    T = BY + H

    # 3 columnas iguales para las tartas
    ncols  = 3
    col_w  = W / ncols
    pad    = col_w * 0.08
    pie_w  = col_w - pad * 2
    pie_h  = H * 0.72
    pie_y  = BY + H * 0.04
    leg_h  = H * 0.22

    for ci, pie_fn in enumerate([_pie_contrato, _pie_jerarquia, _pie_sexo]):
        cx = L + ci * col_w
        # Separador vertical entre columnas
        if ci > 0:
            sep = fig.add_axes([cx - 0.001, BY + 0.02, 0.0015, H - 0.04],
                               facecolor="white", alpha=0.12, zorder=3)
            sep.axis("off")
        # Axes para la tarta
        ax = fig.add_axes([cx + pad, pie_y + leg_h, pie_w, pie_h - leg_h],
                          facecolor="none", zorder=5)
        pie_fn(ax, df)

    _ensure_bg()
    sl = _new_sl(prs); _pic(sl, SHARED_BG, prs)
    _pic(sl, _save_ch(fig, fname), prs)
    _T(sl, slide_title)
    _POP(sl, pop_text)
    _CT(sl, ct_text)
    print(f"  ✓ {slide_title}")


# ── Main ───────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    print("Generando POBLACIONES_v1.pptx (3 slides)…")
    _ensure_bg()
    prs = Presentation()
    prs.slide_width  = Emu(SW_EMU)
    prs.slide_height = Emu(SH_EMU)

    make_slide(
        prs, base,
        slide_title = "Universo Base — Distribución de Población (nº 1.144)",
        pop_text    = "1.144 docentes · Jornada + Honorario · Universo completo de análisis",
        ct_text     = "Distribución de tipo contrato, jerarquía académica y sexo del universo base",
        fname       = "pop_1144.png",
    )
    make_slide(
        prs, ctrl,
        slide_title = "Grupo Control Bloque II — SAT (nº 568)",
        pop_text    = "568 docentes únicos sin formación P3 con SAT disponible (2023–2025)",
        ct_text     = "Distribución de tipo contrato, jerarquía académica y sexo del grupo control B2",
        fname       = "pop_ctrl568.png",
    )
    make_slide(
        prs, sat,
        slide_title = "Aptos P3 — Grupo Tratamiento (nº 210)",
        pop_text    = "210 docentes Aptos P3 con SAT pre+post válido · Grupo de tratamiento",
        ct_text     = "Distribución de tipo contrato, jerarquía académica y sexo de los Aptos P3",
        fname       = "pop_210.png",
    )

    prs.save(OUT_PPTX)
    print(f"\n✓ Guardado: {OUT_PPTX}")
    print(f"  {len(prs.slides)} diapositivas")
